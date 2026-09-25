"""Client delivery files: CSV, Excel (XLSX), a one-page HTML summary, Google Sheets.

Every writer takes one ``DeliveryPackage`` (see ``rows.py``) and uses
``pkg.columns`` for the column order + headers, so all formats show the same
columns in the same order. Row keys that start with ``_`` are internal and are
never written.

Files
-----
``write_all(pkg, folder, formats)`` writes one file per format, named
``<client>-hiring-signals-<YYYY-MM-DD>.<ext>`` (``...-PREVIEW.<ext>`` with
``preview=True``, used for dry runs):

csv   UTF-8 with a BOM so Excel opens accents / emoji correctly; header row =
      the column headers. Formula-injection guard: a text cell that starts with
      ``=`` ``+`` ``-`` ``@`` (or a tab / carriage return) gets a leading ``'``
      so Excel / Sheets show it as text instead of running it as a formula.
xlsx  Sheet "Leads": bold header in the brand colour, frozen header row,
      autofilter, column widths sized to the content (10..60), clickable links
      for Website / Job link / LinkedIn URL, "Date posted" as real dates,
      urgency + email status colour-coded. Text that looks like a formula is
      stored as plain text (quote-prefixed), never as a formula. Sheet "About":
      client, period, counts by signal and by email status, the email-status
      legend and what the columns mean.
html  A branded, self-contained one-pager (inline CSS, no external files unless
      a ``logo_url`` is configured): header, KPI tiles (leads delivered, hot,
      verified emails, companies), leads by signal type, the top-N hottest
      leads and a footer with the sender + website. Print-friendly; every value
      is HTML-escaped; works with 0 rows.

Branding (``pkg.brand`` = the playbook's ``delivery`` section merged with the
client's ``branding``; missing keys fall back to ``playbook.DEFAULTS``):

brand_name    Report name in the HTML header / About sheet (default "Hiring Signal Report").
brand_color   Hex colour like ``#1f4e79`` for headers / accents (invalid -> default).
sender_name   "Prepared by" name (HTML footer, About sheet).
sender_email  Contact email (HTML footer, About sheet).
website       Your website (HTML footer, About sheet).
logo_url      Optional http(s) image shown in the HTML header (with alt text).
footer        Optional extra footer line (HTML footer, About sheet).

Google Sheets
-------------
``push_google_sheet(pkg, cfg, ctx)`` replaces the contents of one worksheet
with the header + rows and returns the sheet URL. Needs the optional
``gspread`` package (``pip install gspread``). Never called in a dry run (it
returns ``""`` if it is). ``cfg`` keys:

spreadsheet_id        The id from the sheet URL (``/spreadsheets/d/<id>/``) - or
spreadsheet_url       the full URL. One of the two is required.
worksheet             Tab name; ``{date}`` (YYYY-MM-DD) and ``{client}`` are
                      replaced (default ``"{date}"``). Created if missing.
service_account_file  Path to the service-account JSON key; default: env
                      ``GOOGLE_APPLICATION_CREDENTIALS``.
service_account_json  Alternatively the key itself (JSON text or mapping), or
                      env ``GOOGLE_SERVICE_ACCOUNT_JSON``.

Values are sent with ``value_input_option=RAW``, so Sheets never evaluates a
cell as a formula and numbers stay numbers.
"""
from __future__ import annotations

import csv
import html
import json
import re
from datetime import date
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple, Union
from urllib.parse import urlparse

from ..outbound.csv_export import FORMULA_PREFIXES, guard_cell
from ..playbook import DEFAULTS
from .rows import EMAIL_LABELS, GUESSED, NOT_FOUND, RISKY, VERIFIED, DeliveryPackage

FORMATS: Tuple[str, ...] = ("csv", "xlsx", "html")

#: Plain-English meaning of every email status (shown in the Excel "About" sheet and the HTML page).
EMAIL_STATUS_MEANINGS: Dict[str, str] = {
    VERIFIED: "Mailbox confirmed deliverable by an email checker or by the data provider "
              "that supplied it. Never a guessed address.",
    RISKY: "A real address from a data provider that could not be confirmed (for example the "
           "company's mail server accepts every address). Usually works; expect some bounces.",
    GUESSED: "Built from the company's usual name pattern (e.g. first.last@company.com). "
             "Not confirmed - check it before you rely on it.",
    NOT_FOUND: "No usable email for this person in this file. Use the LinkedIn URL or the "
               "company website to reach them.",
}

#: Columns explained on the About sheet (only the ones that are not obvious).
COLUMN_NOTES: Dict[str, str] = {
    "signal_type": "Why this company is on the list now (Hiring = a live job posting).",
    "job_titles": "The open role(s) that match your search (up to 5 listed).",
    "posted": "How long ago the job was posted, as of the delivery date.",
    "urgency": "hot = strongest, freshest signals - call these first; normal = good fit, less urgent.",
    "score": "0-100: signal strength + fit with your criteria + how reachable the decision-maker is.",
    "decision_maker": "The person most likely to own this hire (from your buyer titles).",
    "email_status": "How much you can trust the email - see the legend above.",
    "opening_line": "A factual first line you can adapt when you contact the company.",
}

LINK_COLUMNS = frozenset({"website", "job_link", "linkedin_url"})
DATE_COLUMNS = frozenset({"date_posted"})
NUMBER_COLUMNS = frozenset({"score", "company_size"})

MIN_COL_WIDTH = 10
MAX_COL_WIDTH = 60
MAX_LINK_WIDTH = 45
XLSX_MAX_CELL = 32767          # Excel's hard limit per cell
SHEET_TITLE_MAX = 100          # Google Sheets worksheet title limit
DEFAULT_WORKSHEET = "{date}"
JSON_ENV = "GOOGLE_SERVICE_ACCOUNT_JSON"
FILE_ENV = "GOOGLE_APPLICATION_CREDENTIALS"

_HEX_RE = re.compile(r"^#[0-9A-Fa-f]{6}$")
_BARE_DOMAIN_RE = re.compile(r"^(?:www\.)?[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)*\.[A-Za-z]{2,}(?:[/?#]\S*)?$")
# Control characters Excel files cannot hold (openpyxl raises IllegalCharacterError on them).
_XLSX_ILLEGAL_RE = re.compile(r"[\000-\010\013\014\016-\037]")
_SAFE_NAME_RE = re.compile(r"[^A-Za-z0-9_-]+")


# --- shared helpers ------------------------------------------------------------------

def brand_settings(pkg: DeliveryPackage) -> Dict[str, str]:
    """``pkg.brand`` over the playbook defaults, as clean strings (invalid colour -> default)."""
    defaults = dict(DEFAULTS.get("delivery") or {})
    out: Dict[str, str] = {}
    brand = pkg.brand if isinstance(pkg.brand, Mapping) else {}
    for key in set(defaults) | set(brand):
        value = brand.get(key)
        text = "" if value is None else str(value).strip()
        out[key] = text or str(defaults.get(key) or "")
    out["brand_name"] = out.get("brand_name") or "Hiring Signal Report"
    color = out.get("brand_color", "")
    if not _HEX_RE.match(color):
        color = str(defaults.get("brand_color") or "#1f4e79")
    out["brand_color"] = color.lower()
    return out


def _luminance(hex_color: str) -> float:
    """WCAG relative luminance of ``#rrggbb`` (0 = black, 1 = white)."""
    h = hex_color.lstrip("#")
    chans = []
    for i in (0, 2, 4):
        c = int(h[i:i + 2], 16) / 255
        chans.append(c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4)
    r, g, b = chans
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def ink_on(hex_color: str) -> str:
    """Readable text colour (white or near-black) on a background of ``hex_color``."""
    lum = _luminance(hex_color)
    # contrast against white vs against #1c2430 (luminance ~0.017)
    return "#ffffff" if (1.05 / (lum + 0.05)) >= ((lum + 0.05) / 0.067) else "#1c2430"


def _text(value: Any) -> str:
    return "" if value is None else str(value)


def link_target(value: Any) -> str:
    """The http(s) URL to link to for a cell value, or '' when it is not a safe web link.

    ``https://...`` / ``http://...`` are used as-is; a bare domain like ``acme.com`` or
    ``linkedin.com/in/jane`` gets ``https://``. Anything else (``javascript:``, text) -> ''.
    """
    text = _text(value).strip()
    if not text or any(ch.isspace() for ch in text):
        return ""
    parsed = urlparse(text)
    if parsed.scheme.lower() in ("http", "https"):
        return text if parsed.netloc else ""
    if not parsed.scheme and _BARE_DOMAIN_RE.match(text):
        return "https://" + text
    return ""


def _parse_iso_date(value: Any) -> Optional[date]:
    if isinstance(value, date):
        return value
    text = _text(value).strip()
    if not text:
        return None
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        return None


def companies_count(pkg: DeliveryPackage) -> int:
    """Distinct companies in the package (by internal company key, else by name)."""
    keys = set()
    for r in pkg.rows:
        key = r.get("_company_key") or _text(r.get("company")).strip().lower()
        if key:
            keys.add(key)
    return len(keys)


def file_name(pkg: DeliveryPackage, ext: str, preview: bool = False) -> str:
    """``<client>-hiring-signals-<YYYY-MM-DD>[-PREVIEW].<ext>`` (client name made file-safe)."""
    client = _SAFE_NAME_RE.sub("-", _text(pkg.client_name).strip()).strip("-") or "client"
    day = pkg.period_end.isoformat() if isinstance(pkg.period_end, date) else _text(pkg.period_end)
    return f"{client}-hiring-signals-{day}{'-PREVIEW' if preview else ''}.{ext.lstrip('.')}"


def _period_text(pkg: DeliveryPackage) -> str:
    start, end = pkg.period_start, pkg.period_end
    if start and end and start != end:
        return f"{start.isoformat()} to {end.isoformat()}"
    return (end or start).isoformat() if (end or start) else ""


# --- CSV ------------------------------------------------------------------------------

def write_csv(pkg: DeliveryPackage, path: Union[str, Path]) -> Path:
    """Write the leads as a CSV (UTF-8 with BOM, formula-guarded). Returns the path."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    header = [h for _, h in pkg.columns]
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(header)
        for row in pkg.public_rows():
            writer.writerow([guard_cell("" if v is None else v) for v in row])
    return path


# --- XLSX -----------------------------------------------------------------------------

def _xlsx_text(value: Any) -> str:
    """Text safe for an Excel cell: no illegal control characters, within the size limit."""
    text = _XLSX_ILLEGAL_RE.sub("", _text(value))
    return text[:XLSX_MAX_CELL]


def _put(ws: Any, row: int, col: int, value: Any) -> Any:
    """Write one value; strings that look like formulas are stored as quote-prefixed text."""
    if value is None or value == "":
        return ws.cell(row=row, column=col, value=None)
    if isinstance(value, bool) or not isinstance(value, (int, float, date)):
        value = _xlsx_text(value)
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str) and value[:1] in FORMULA_PREFIXES:
        cell.data_type = "s"          # plain text, never a formula
        cell.quotePrefix = True       # like typing a leading ' in Excel
    return cell


def _column_width(header: str, values: Iterable[Any], key: str) -> float:
    longest = len(header)
    for v in values:
        if isinstance(v, date):
            n = 10
        else:
            text = _text(v)
            n = max((len(line) for line in text.splitlines()), default=0)
        longest = max(longest, n)
    cap = MAX_LINK_WIDTH if key in LINK_COLUMNS else MAX_COL_WIDTH
    return float(max(MIN_COL_WIDTH, min(cap, longest + 2)))


def write_xlsx(pkg: DeliveryPackage, path: Union[str, Path]) -> Path:
    """Write the Excel workbook (sheets "Leads" + "About"). Returns the path."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as e:  # pragma: no cover - openpyxl is a hard dependency
        raise RuntimeError("Excel output needs the openpyxl package: pip install openpyxl") from e

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    brand = brand_settings(pkg)
    color = brand["brand_color"].lstrip("#").upper()
    header_font = Font(bold=True, color="FF" + ink_on(brand["brand_color"]).lstrip("#").upper())
    header_fill = PatternFill(fill_type="solid", start_color="FF" + color, end_color="FF" + color)
    link_font = Font(color="FF0563C1", underline="single")

    status_fills = {
        VERIFIED: PatternFill(fill_type="solid", start_color="FFE6F4EA", end_color="FFE6F4EA"),
        RISKY: PatternFill(fill_type="solid", start_color="FFFFF4E5", end_color="FFFFF4E5"),
        GUESSED: PatternFill(fill_type="solid", start_color="FFF1EEFB", end_color="FFF1EEFB"),
        NOT_FOUND: PatternFill(fill_type="solid", start_color="FFF2F4F7", end_color="FFF2F4F7"),
    }
    hot_fill = PatternFill(fill_type="solid", start_color="FFFDECEA", end_color="FFFDECEA")
    hot_font = Font(bold=True, color="FFB42318")

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    cols = pkg.columns
    keys = [k for k, _ in cols]

    for ci, (_, header) in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=ci, value=header)
        cell.font, cell.fill = header_font, header_fill
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 20

    for ri, row in enumerate(pkg.rows, start=2):
        for ci, key in enumerate(keys, start=1):
            value = row.get(key, "")
            if key in DATE_COLUMNS:
                day = _parse_iso_date(value)
                if day is not None:
                    cell = ws.cell(row=ri, column=ci, value=day)
                    cell.number_format = "yyyy-mm-dd"
                    continue
            if key in NUMBER_COLUMNS and isinstance(value, (int, float)) and not isinstance(value, bool):
                cell = ws.cell(row=ri, column=ci, value=value)
                cell.number_format = "#,##0" if key == "company_size" else "0"
                continue
            cell = _put(ws, ri, ci, value)
            if key in LINK_COLUMNS:
                target = link_target(value)
                if target:
                    cell.hyperlink = target
                    cell.font = link_font
            elif key == "email_status" and value in status_fills:
                cell.fill = status_fills[value]
            elif key == "urgency" and value == "hot":
                cell.fill, cell.font = hot_fill, hot_font

    last_col = get_column_letter(max(1, len(cols)))
    last_row = max(1, len(pkg.rows) + 1)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{last_col}{last_row}"
    for ci, (key, header) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = _column_width(
            header, (r.get(key, "") for r in pkg.rows[:500]), key)
    ws.sheet_properties.tabColor = color
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"

    _write_about(wb.create_sheet("About"), pkg, brand)

    wb.properties.title = f"{brand['brand_name']} - {pkg.client_display}"
    wb.properties.creator = brand.get("sender_name") or brand["brand_name"]
    wb.save(path)
    return path


def _write_about(ws: Any, pkg: DeliveryPackage, brand: Dict[str, str]) -> None:
    """Fill the "About" sheet: who / when / counts / email-status legend / column notes."""
    from openpyxl.styles import Alignment, Font, PatternFill

    color = brand["brand_color"].lstrip("#").upper()
    title_font = Font(bold=True, size=16, color="FF" + (color if _luminance(brand["brand_color"]) < 0.4
                                                         else "1C2430"))
    section_font = Font(bold=True, size=12)
    bold = Font(bold=True)
    section_fill = PatternFill(fill_type="solid", start_color="FFF2F4F7", end_color="FFF2F4F7")
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")
    state = {"row": 1}

    def line(label: Any = "", value: Any = "", *, label_font: Any = None) -> None:
        r = state["row"]
        a = _put(ws, r, 1, label)
        b = _put(ws, r, 2, value)
        a.alignment, b.alignment = top, wrap
        if label_font is not None:
            a.font = label_font
        state["row"] += 1

    def section(title: str, head: Tuple[str, str]) -> None:
        state["row"] += 1
        r = state["row"]
        cell = _put(ws, r, 1, title)
        cell.font = section_font
        for ci in (1, 2):
            ws.cell(row=r, column=ci).fill = section_fill
        state["row"] += 1
        line(head[0], head[1])
        for ci in (1, 2):
            ws.cell(row=state["row"] - 1, column=ci).font = bold

    cell = _put(ws, 1, 1, brand["brand_name"])
    cell.font = title_font
    state["row"] = 3
    line("Prepared for", pkg.client_display or pkg.client_name, label_font=bold)
    line("Period", _period_text(pkg), label_font=bold)
    line("Delivered on", pkg.period_end.isoformat() if pkg.period_end else "", label_font=bold)
    line("Leads in this file", len(pkg.rows), label_font=bold)
    line("Hot leads", pkg.hot, label_font=bold)
    line("Companies", companies_count(pkg), label_font=bold)
    for i, note in enumerate(pkg.notes or []):
        line("Notes" if i == 0 else "", note, label_font=bold)

    section("Leads by signal", ("Signal", "Leads"))
    if pkg.counts_by_signal:
        for label, n in pkg.counts_by_signal.items():
            line(label, n)
    else:
        line("(none)", 0)

    section("Leads by email status", ("Email status", "Leads"))
    counts = pkg.counts_by_email_status
    for status in EMAIL_LABELS:
        line(status, counts.get(status, 0))

    section("Email status legend", ("Email status", "What it means"))
    for status in EMAIL_LABELS:
        line(status, EMAIL_STATUS_MEANINGS[status])

    section("Columns", ("Column", "What it means"))
    headers = dict(pkg.columns)
    for key, meaning in COLUMN_NOTES.items():
        if key in headers:
            line(headers[key], meaning)

    who = [x for x in (brand.get("sender_name"), brand.get("sender_email"), brand.get("website")) if x]
    if who or brand.get("footer"):
        state["row"] += 1
        if who:
            line("Prepared by", " | ".join(who), label_font=bold)
        if brand.get("footer"):
            line("", brand["footer"])

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 90
    ws.sheet_view.showGridLines = False


# --- HTML -----------------------------------------------------------------------------

def _e(value: Any) -> str:
    return html.escape(_text(value), quote=True)


def _day(d: Optional[date]) -> str:
    return f"{d.day} {d.strftime('%b %Y')}" if isinstance(d, date) else ""


def _posted_short(row: Mapping[str, Any]) -> str:
    posted = _text(row.get("posted")).strip()
    if posted.lower().startswith("posted "):
        posted = posted[7:]
    return posted or _text(row.get("date_posted"))


def _link(text: str, url: Any, cls: str = "") -> str:
    target = link_target(url)
    if not target:
        return _e(text)
    attr = f' class="{cls}"' if cls else ""
    return f'<a href="{_e(target)}"{attr} rel="noopener noreferrer">{_e(text)}</a>'


def _hot_first(rows: Sequence[Mapping[str, Any]]) -> List[Mapping[str, Any]]:
    """Rows in their given order, with hot leads moved to the front (stable)."""
    return sorted(rows, key=lambda r: 0 if r.get("urgency") == "hot" else 1)


_STATUS_CLASS = {VERIFIED: "st-verified", RISKY: "st-risky", GUESSED: "st-guessed", NOT_FOUND: "st-none"}

_CSS = """
:root{--brand:%(brand)s;--brand-ink:%(ink)s;--brand-text:%(brand_text)s;--ink:#1c2430;--muted:#5b6573;
--line:#e3e7ec;--bg:#f4f6f9;--card:#ffffff}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--ink);
font:14px/1.5 -apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,"Helvetica Neue",Arial,sans-serif}
.page{max-width:1040px;margin:24px auto;background:var(--card);border-radius:10px;overflow:hidden;
box-shadow:0 1px 3px rgba(16,24,40,.08),0 1px 2px rgba(16,24,40,.04)}
.masthead{background:var(--brand);color:var(--brand-ink);padding:28px 32px;display:flex;gap:20px;align-items:center}
.masthead img{max-height:52px;max-width:180px;display:block}
.eyebrow{margin:0;font-size:12px;letter-spacing:.08em;text-transform:uppercase;opacity:.85}
h1{margin:4px 0 2px;font-size:24px;line-height:1.25}
.period{margin:0;opacity:.9}
.content{padding:24px 32px 8px}
.kpis{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:12px;margin:0 0 8px}
.kpi{border:1px solid var(--line);border-radius:8px;padding:14px 16px}
.kpi .value{display:block;font-size:28px;font-weight:700;line-height:1.2;color:var(--brand-text);
font-variant-numeric:tabular-nums}
.kpi .label{display:block;color:var(--muted);font-size:12px;letter-spacing:.04em;text-transform:uppercase}
h2{margin:28px 0 10px;font-size:16px}
table{width:100%%;border-collapse:collapse}
th{text-align:left;font-size:12px;font-weight:600;color:var(--muted);letter-spacing:.04em;text-transform:uppercase;
border-bottom:2px solid var(--line);padding:8px 10px}
td{border-bottom:1px solid var(--line);padding:9px 10px;vertical-align:top}
td.num,th.num{text-align:right;font-variant-numeric:tabular-nums;white-space:nowrap}
.signals td.bar{width:45%%}
.bar-track{background:#eef1f5;height:8px;border-radius:4px;margin-top:7px}
.bar-fill{background:var(--brand);height:8px;border-radius:4px}
.sub{display:block;color:var(--muted);font-size:12px}
.badge{display:inline-block;padding:1px 8px;border-radius:999px;font-size:12px;font-weight:600;white-space:nowrap}
.u-hot{background:#fdecea;color:#b42318}
.u-normal{background:#eef1f5;color:#344054}
.st-verified{background:#e6f4ea;color:#067647}
.st-risky{background:#fff4e5;color:#b54708}
.st-guessed{background:#f1eefb;color:#5b3fb0}
.st-none{background:#f2f4f7;color:#475467}
a{color:var(--brand-text)}
.note{color:var(--muted);font-size:13px;margin:10px 0 0}
.empty{padding:28px;border:1px dashed #cfd6de;border-radius:8px;text-align:center;color:var(--muted)}
.empty strong{display:block;color:var(--ink);font-size:15px;margin-bottom:4px}
.legend{margin:8px 0 0;padding:0;list-style:none;color:var(--muted);font-size:13px}
.legend li{margin:4px 0}
.notes{margin:8px 0 0;padding-left:18px;color:var(--muted);font-size:13px}
footer{margin-top:24px;padding:16px 32px 24px;border-top:1px solid var(--line);color:var(--muted);font-size:12px}
footer p{margin:2px 0}
@media (max-width:720px){.kpis{grid-template-columns:repeat(2,minmax(0,1fr))}
.masthead,.content,footer{padding-left:16px;padding-right:16px}
.table-wrap{overflow-x:auto}}
@media print{body{background:#fff;font-size:12px}
.page{max-width:none;margin:0;border-radius:0;box-shadow:none}
.masthead,.badge,.bar-fill,.bar-track{-webkit-print-color-adjust:exact;print-color-adjust:exact}
a{color:inherit;text-decoration:none}
tr,.kpi,.empty{break-inside:avoid;page-break-inside:avoid}
thead{display:table-header-group}
@page{margin:12mm}}
"""


def render_html(pkg: DeliveryPackage, top: int = 10) -> str:
    """The branded one-page summary as an HTML string (see module docstring)."""
    brand = brand_settings(pkg)
    color = brand["brand_color"]
    brand_text = color if (1.05 / (_luminance(color) + 0.05)) >= 3.0 else "#1c2430"
    css = _CSS % {"brand": color, "ink": ink_on(color), "brand_text": brand_text}
    total = len(pkg.rows)
    counts = pkg.counts_by_email_status
    client = pkg.client_display or pkg.client_name
    start, end = _day(pkg.period_start), _day(pkg.period_end)
    period = f"{start} &ndash; {end}" if start and end and start != end else (end or start)
    title = " - ".join(x for x in (brand["brand_name"], client, pkg.period_end.isoformat()
                                    if pkg.period_end else "") if x)

    parts: List[str] = [
        "<!DOCTYPE html>",
        '<html lang="en">',
        "<head>",
        '<meta charset="utf-8">',
        '<meta name="viewport" content="width=device-width, initial-scale=1">',
        f"<title>{_e(title)}</title>",
        f"<style>{css}</style>",
        "</head>",
        "<body>",
        '<main class="page">',
        '<header class="masthead">',
    ]
    logo = brand.get("logo_url", "")
    # only a full http(s) URL is used (no bare domains, no javascript: / data: tricks)
    if logo.lower().startswith(("http://", "https://")) and link_target(logo):
        parts.append(f'<img src="{_e(logo)}" alt="{_e(brand["brand_name"])} logo">')
    parts += [
        "<div>",
        f'<p class="eyebrow">{_e(brand["brand_name"])}</p>',
        f"<h1>Prepared for {_e(client)}</h1>",
        f'<p class="period">{period}</p>' if period else "",
        "</div>",
        "</header>",
        '<div class="content">',
        '<section class="kpis" aria-label="Key numbers">',
    ]
    for value, label in ((total, "Leads delivered"), (pkg.hot, "Hot leads"),
                         (counts.get(VERIFIED, 0), "Verified emails"), (companies_count(pkg), "Companies")):
        parts.append(f'<div class="kpi"><span class="value">{_e(value)}</span>'
                     f'<span class="label">{_e(label)}</span></div>')
    parts.append("</section>")

    if not total:
        parts += [
            '<div class="empty">',
            "<strong>No new leads matched your criteria this period.</strong>",
            "Nothing new was posted that fits your roles, locations and company size, or "
            "everything that did was already delivered to you. We keep looking every week.",
            "</div>",
        ]
    else:
        parts += _signals_section(pkg)
        parts += _top_section(pkg, top)

    if pkg.notes:
        parts.append("<h2>About this report</h2>")
        parts.append('<ul class="notes">' + "".join(f"<li>{_e(n)}</li>" for n in pkg.notes) + "</ul>")

    parts.append("</div>")
    parts += _footer(pkg, brand)
    parts += ["</main>", "</body>", "</html>", ""]
    return "\n".join(p for p in parts if p != "") + "\n"


def _signals_section(pkg: DeliveryPackage) -> List[str]:
    counts = pkg.counts_by_signal
    biggest = max(counts.values()) if counts else 1
    total = len(pkg.rows) or 1
    out = ["<h2>Leads by signal type</h2>",
           '<table class="signals"><thead><tr><th>Signal</th><th class="num">Leads</th>'
           '<th class="num">Share</th><th></th></tr></thead><tbody>']
    for label, n in counts.items():
        width = max(2, round(100 * n / biggest))
        out.append(f"<tr><td>{_e(label)}</td><td class=\"num\">{n}</td>"
                   f"<td class=\"num\">{round(100 * n / total)}%</td>"
                   f'<td class="bar"><div class="bar-track"><div class="bar-fill" '
                   f'style="width:{width}%"></div></div></td></tr>')
    out.append("</tbody></table>")
    return out


def _top_section(pkg: DeliveryPackage, top: int) -> List[str]:
    top = max(0, int(top or 0))
    picked = _hot_first(pkg.rows)[:top]
    if not picked:
        return []
    if len(pkg.rows) > len(picked):
        heading = f"Top {len(picked)} hottest leads"
    else:
        heading = "Your lead" if len(picked) == 1 else f"All {len(picked)} leads, hottest first"
    out = [f"<h2>{_e(heading)}</h2>", '<div class="table-wrap"><table class="leads"><thead><tr>'
           "<th>Company</th><th>Role(s)</th><th>Posted</th><th>Location</th><th>Urgency</th>"
           "<th>Decision-maker</th><th>Email status</th></tr></thead><tbody>"]
    for r in picked:
        company = _link(_text(r.get("company")) or "(unnamed company)", r.get("website"))
        roles = _text(r.get("job_titles")) or _text(r.get("signal_type"))
        role_html = _link(roles, r.get("job_link")) if roles else '<span class="sub">-</span>'
        sig = _text(r.get("signal_type"))
        if sig and sig != "Hiring" and roles != sig:
            role_html += f'<span class="sub">{_e(sig)}</span>'
        urgency = _text(r.get("urgency"))
        u_cls = "u-hot" if urgency == "hot" else "u-normal"
        name, dm_title = _text(r.get("decision_maker")), _text(r.get("decision_maker_title"))
        dm = _e(name) if name else '<span class="sub">not identified</span>'
        if name and dm_title:
            dm += f'<span class="sub">{_e(dm_title)}</span>'
        status = _text(r.get("email_status")) or NOT_FOUND
        out.append(
            "<tr>"
            f"<td>{company}</td>"
            f"<td>{role_html}</td>"
            f"<td>{_e(_posted_short(r))}</td>"
            f"<td>{_e(r.get('location'))}</td>"
            f'<td><span class="badge {u_cls}">{_e(urgency)}</span></td>'
            f"<td>{dm}</td>"
            f'<td><span class="badge {_STATUS_CLASS.get(status, "st-none")}">{_e(status)}</span></td>'
            "</tr>")
    out.append("</tbody></table></div>")
    if len(pkg.rows) > len(picked):
        out.append(f'<p class="note">Showing {len(picked)} of {len(pkg.rows)} leads. The full list, '
                   "with emails and LinkedIn profiles, is in the Excel / CSV file.</p>")
    else:
        out.append('<p class="note">Emails and LinkedIn profiles are in the Excel / CSV file.</p>')
    out.append('<ul class="legend">' + "".join(
        f'<li><span class="badge {_STATUS_CLASS[s]}">{_e(s)}</span> {_e(EMAIL_STATUS_MEANINGS[s])}</li>'
        for s in EMAIL_LABELS) + "</ul>")
    return out


def _footer(pkg: DeliveryPackage, brand: Dict[str, str]) -> List[str]:
    bits: List[str] = []
    if brand.get("sender_name"):
        bits.append(f"Prepared by {_e(brand['sender_name'])}")
    email = brand.get("sender_email", "")
    if email and "@" in email and not any(ch.isspace() for ch in email):
        bits.append(f'<a href="mailto:{_e(email)}">{_e(email)}</a>')
    elif email:
        bits.append(_e(email))
    site = brand.get("website", "")
    if site:
        shown = re.sub(r"^https?://", "", site, flags=re.I).rstrip("/") or site
        bits.append(_link(shown, site))
    out = ["<footer>"]
    if bits:
        out.append("<p>" + " &middot; ".join(bits) + "</p>")
    if brand.get("footer"):
        out.append(f"<p>{_e(brand['footer'])}</p>")
    generated = pkg.period_end.isoformat() if pkg.period_end else ""
    out.append(f"<p>{_e(brand['brand_name'])}{' - ' + _e(generated) if generated else ''}</p>")
    out.append("</footer>")
    return out


def write_html(pkg: DeliveryPackage, path: Union[str, Path], top: int = 10) -> Path:
    """Write the one-page HTML summary (UTF-8). Returns the path."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(render_html(pkg, top=top), encoding="utf-8")
    return path


# --- write_all ------------------------------------------------------------------------

WRITERS: Dict[str, Callable[[DeliveryPackage, Path], Path]] = {
    "csv": write_csv,
    "xlsx": write_xlsx,
    "html": write_html,
}


def normalize_formats(formats: Union[str, Iterable[Any], None]) -> List[str]:
    """Lower-cased, de-duplicated format names; ValueError naming any unknown one."""
    if formats is None:
        return []
    items = [formats] if isinstance(formats, str) else list(formats)
    out: List[str] = []
    bad: List[str] = []
    for f in items:
        name = _text(f).strip().lower().lstrip(".")
        if not name:
            continue
        if name not in WRITERS:
            bad.append(_text(f))
        elif name not in out:
            out.append(name)
    if bad:
        raise ValueError(f"unknown delivery format(s): {', '.join(repr(b) for b in bad)} - "
                         f"choose from {', '.join(FORMATS)}")
    return out


def write_all(pkg: DeliveryPackage, folder: Union[str, Path], formats: Union[str, Iterable[Any]],
              *, preview: bool = False) -> Dict[str, Path]:
    """Write every requested format into ``folder``; returns ``{format: path}`` in request order."""
    names = normalize_formats(formats)
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    return {name: WRITERS[name](pkg, folder / file_name(pkg, name, preview)) for name in names}


# --- Google Sheets --------------------------------------------------------------------

class SheetPushError(RuntimeError):
    """The Google Sheet could not be opened or written."""


def _import_gspread() -> Any:
    try:
        import gspread  # type: ignore[import-not-found]
    except ImportError as e:
        raise RuntimeError("Pushing to Google Sheets needs the gspread package - install it with: "
                           "pip install gspread   (or leave delivery.google_sheet.spreadsheet_id empty)") from e
    return gspread


def _is_exc(exc: BaseException, gspread: Any, name: str) -> bool:
    """isinstance check against ``gspread.exceptions.<name>`` that tolerates missing classes."""
    cls = getattr(getattr(gspread, "exceptions", None), name, None) or getattr(gspread, name, None)
    if isinstance(cls, type) and isinstance(exc, cls):
        return True
    return type(exc).__name__ == name


def _sheet_ref(cfg: Mapping[str, Any]) -> Tuple[str, str]:
    sid = _text(cfg.get("spreadsheet_id")).strip()
    url = _text(cfg.get("spreadsheet_url")).strip()
    if sid.startswith(("http://", "https://")):
        sid, url = "", url or sid
    if sid:
        return "key", sid
    if url:
        return "url", url
    raise ValueError("Google Sheets push: set google_sheet.spreadsheet_id (the id in the sheet's URL, "
                     ".../spreadsheets/d/<id>/) or google_sheet.spreadsheet_url")


def worksheet_title(template: Any, pkg: DeliveryPackage) -> str:
    """Worksheet name from the template (``{date}`` / ``{client}``), max 100 characters."""
    text = _text(template).strip() or DEFAULT_WORKSHEET
    day = pkg.period_end.isoformat() if pkg.period_end else ""
    text = text.replace("{date}", day).replace("{client}", _text(pkg.client_name))
    text = re.sub(r"\s+", " ", text).strip()
    return text[:SHEET_TITLE_MAX] or day or "leads"


def _env(ctx: Any) -> Mapping[str, str]:
    env = getattr(ctx, "env", None)
    return env if isinstance(env, Mapping) else {}


def _sheets_client(gspread: Any, cfg: Mapping[str, Any], ctx: Any) -> Any:
    raw = cfg.get("service_account_json") or _env(ctx).get(JSON_ENV)
    if raw:
        if isinstance(raw, Mapping):
            info: Any = dict(raw)
        else:
            try:
                info = json.loads(_text(raw))
            except ValueError as e:
                raise SheetPushError(f"Google Sheets push: the service account JSON is not valid JSON ({e})") from None
        if not isinstance(info, dict):
            raise SheetPushError("Google Sheets push: the service account JSON must be a JSON object")
        return gspread.service_account_from_dict(info)
    key_path = _text(cfg.get("service_account_file")).strip() or _text(_env(ctx).get(FILE_ENV)).strip()
    if not key_path:
        raise SheetPushError(
            "Google Sheets push: no Google credentials - set google_sheet.service_account_file (path to a "
            f"service-account JSON key) or the {FILE_ENV} environment variable")
    playbook = getattr(ctx, "playbook", None)
    resolved = playbook.resolve_path(key_path) if hasattr(playbook, "resolve_path") else Path(key_path)
    if not Path(resolved).exists():
        raise SheetPushError(f"Google Sheets push: service account key file not found: {resolved}")
    return gspread.service_account(filename=str(resolved))


def _sheet_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float, str)):
        return value
    return _text(value)


def push_google_sheet(pkg: DeliveryPackage, cfg: Optional[Mapping[str, Any]], ctx: Any) -> str:
    """Replace one worksheet's contents with the delivery rows; returns the sheet URL.

    Returns ``""`` without touching Google in a dry run. Raises ``ValueError`` for missing
    config, ``RuntimeError`` (with an install hint) when gspread is missing and
    ``SheetPushError`` when the sheet cannot be opened or written.
    """
    cfg = dict(cfg or {})
    log = getattr(ctx, "log", None)
    title = worksheet_title(cfg.get("worksheet"), pkg)
    if getattr(ctx, "dry_run", False):
        if log is not None:
            log.info("google sheet: dry run - not writing %d row(s) to worksheet '%s'", len(pkg.rows), title)
        return ""
    kind, ref = _sheet_ref(cfg)
    gspread = _import_gspread()
    client = _sheets_client(gspread, cfg, ctx)
    try:
        book = client.open_by_key(ref) if kind == "key" else client.open_by_url(ref)
    except Exception as e:  # noqa: BLE001 - translate gspread errors into one clear message
        if _is_exc(e, gspread, "SpreadsheetNotFound") or _is_exc(e, gspread, "NoValidUrlKeyFound"):
            raise SheetPushError(f"Google Sheets push: spreadsheet {ref!r} not found - check the id and share "
                                 "the sheet with the service account's email as Editor") from e
        raise SheetPushError(f"Google Sheets push: could not open spreadsheet {ref!r} ({e}) - make sure the "
                             "Google Sheets API is enabled and the sheet is shared with the service account") from e

    header = [h for _, h in pkg.columns]
    values = [header] + [[_sheet_value(v) for v in row] for row in pkg.public_rows()]
    try:
        try:
            ws = book.worksheet(title)
        except Exception as e:  # noqa: BLE001
            if not _is_exc(e, gspread, "WorksheetNotFound"):
                raise
            ws = book.add_worksheet(title=title, rows=max(len(values), 100), cols=max(len(header), 26))
        ws.clear()
        try:
            cur_rows, cur_cols = int(ws.row_count), int(ws.col_count)
        except (AttributeError, TypeError, ValueError):
            cur_rows, cur_cols = len(values), len(header)
        if cur_rows < len(values) or cur_cols < len(header):
            ws.resize(rows=max(cur_rows, len(values)), cols=max(cur_cols, len(header)))
        ws.update(range_name="A1", values=values, value_input_option="RAW")
    except SheetPushError:
        raise
    except Exception as e:  # noqa: BLE001
        raise SheetPushError(f"Google Sheets push: could not write worksheet '{title}' ({e})") from e
    for cosmetic in (lambda: ws.freeze(rows=1),
                     lambda: ws.format("1:1", {"textFormat": {"bold": True}})):
        try:
            cosmetic()
        except Exception as e:  # noqa: BLE001 - formatting is optional
            if log is not None:
                log.debug("google sheet: formatting skipped: %s", e)

    url = _text(getattr(book, "url", "")).strip()
    if not url:
        url = f"https://docs.google.com/spreadsheets/d/{getattr(book, 'id', '') or ref}"
    gid = getattr(ws, "id", None)
    if isinstance(gid, int) and not isinstance(gid, bool) and "#gid=" not in url:
        url = f"{url}#gid={gid}"
    if log is not None:
        log.info("google sheet: worksheet '%s' replaced with %d row(s) (%s)", title, len(pkg.rows), url)
    return url


__all__ = [
    "EMAIL_STATUS_MEANINGS", "FORMATS", "SheetPushError", "brand_settings", "companies_count", "file_name",
    "link_target", "normalize_formats", "push_google_sheet", "render_html", "worksheet_title", "write_all",
    "write_csv", "write_html", "write_xlsx",
]
