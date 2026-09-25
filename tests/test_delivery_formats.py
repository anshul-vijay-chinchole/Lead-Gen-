"""Tests for leadgen.delivery.formats: CSV / XLSX / HTML writers, write_all, Google Sheets push."""
from __future__ import annotations

import csv
import re
import sys
import types
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace
from typing import Any, Dict, List, Optional

import openpyxl
import pytest

from leadgen.delivery import formats
from leadgen.delivery.formats import (EMAIL_STATUS_MEANINGS, SheetPushError, brand_settings, file_name,
                                      link_target, push_google_sheet, render_html, write_all, write_csv,
                                      write_html, write_xlsx)
from leadgen.delivery.rows import (EMAIL_LABELS, GUESSED, NOT_FOUND, RISKY, VERIFIED, DeliveryPackage,
                                   build_row)
from leadgen.models import Company, Contact, EmailStatus, Lead, Signal
from tests.conftest import TODAY

START = date(2026, 9, 17)


# --- builders ---------------------------------------------------------------------------

def make_lead(name: str = "Acme Corp", domain: str = "acme.com", *, title: str = "Senior Accountant",
              posted: Optional[date] = date(2026, 9, 22), location: str = "Austin, TX",
              tier: str = "hot", score: int = 88, employees: Optional[int] = 120,
              industry: str = "Accounting", url: str = "https://jobs.acme.com/123",
              signal_type: str = "job_posting", website: str = "",
              contact: Any = "default", source: str = "theirstack") -> Lead:
    sig = Signal(type=signal_type, title=title, source=source, posted_at=posted, url=url,
                 location=location, external_id=f"{name}-1")
    company = Company(name=name, domain=domain, website=website, location=location, industry=industry,
                      employees=employees, signals=[sig] if title or url else [])
    if contact == "default":
        contact = Contact(first_name="Jane", last_name="Doe", title="CFO", email=f"jane@{domain or 'x.com'}",
                          email_status=EmailStatus.VALID, source="apollo",
                          linkedin_url="https://www.linkedin.com/in/jane-doe")
    return Lead(company=company, contact=contact, score=score, tier=tier, playbook="client-acme")


def tricky_leads() -> List[Lead]:
    """Unicode, CSV-hostile text, formula-looking values, missing fields, every email status."""
    return [
        make_lead(),  # verified email
        make_lead("Société Générale Ünïcødé 株式会社 🚀", "societe.example", title="Comptable — Senior (H/F)",
                  location="Paris, Île-de-France", tier="hot", score=91,
                  contact=Contact(full_name="Zoë Ångström", title="DAF", email="zoe@societe.example",
                                  email_status=EmailStatus.UNKNOWN, source="hunter")),        # risky
        make_lead('Smith, "Jones" & Co\nLtd', "smithjones.example", title='Controller, "Group"\nFinance',
                  tier="normal", score=70,
                  contact=Contact(full_name="Pat Smith", title="VP Finance", email="pat.smith@smithjones.example",
                                  email_status=EmailStatus.VALID, source="pattern")),          # guessed
        make_lead('=HYPERLINK("http://evil.example","click")', "evil.example", title="+cmd|' /C calc'!A0",
                  location="-2+3", tier="normal", score=65, contact=None),                     # not found
        make_lead("@SUM Industries", "", title="Accountant\x07 (bell)", posted=None, location="",
                  employees=None, industry="", url="", website="javascript:alert(1)", tier="normal", score=61,
                  contact=Contact(full_name="Lee Wong", title="Controller", email="",
                                  linkedin_url="linkedin.com/in/lee-wong")),                  # not found
    ]


def make_pkg(leads: List[Lead], *, include_opening: bool = False, brand: Optional[Dict[str, Any]] = None,
             notes: Optional[List[str]] = None, client: str = "acme",
             display: str = "Acme Staffing Ltd") -> DeliveryPackage:
    rows = [build_row(ld, TODAY, include_opening=include_opening,
                      opening_line=f"Saw {ld.company.name} is hiring." if include_opening else "")
            for ld in leads]
    return DeliveryPackage(client_name=client, client_display=display, period_start=START, period_end=TODAY,
                           rows=rows, include_opening=include_opening, brand=dict(brand or {}),
                           run_id="run-1", notes=list(notes or []))


def many_leads(n: int) -> List[Lead]:
    out = []
    for i in range(n):
        out.append(make_lead(f"Company {i:02d}", f"company{i:02d}.example", title=f"Accountant {i}",
                             tier="hot" if i < n // 3 else "normal", score=95 - (i % 30)))
    return out


def read_csv(path: Path) -> List[List[str]]:
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.reader(f))


# --- CSV --------------------------------------------------------------------------------

def test_csv_header_bom_roundtrip_and_injection_guard(tmp_path):
    pkg = make_pkg(tricky_leads())
    path = write_csv(pkg, tmp_path / "out.csv")
    assert path == tmp_path / "out.csv"
    assert path.read_bytes().startswith(b"\xef\xbb\xbf")
    rows = read_csv(path)
    header = [h for _, h in pkg.columns]
    assert rows[0] == header
    assert len(rows) == 1 + len(pkg.rows)
    by_company = {r[0]: dict(zip(header, r)) for r in rows[1:]}

    # unicode and CSV-hostile text survive the round trip
    assert "Société Générale Ünïcødé 株式会社 🚀" in by_company
    tricky = by_company['Smith, "Jones" & Co\nLtd']
    assert tricky["Job title(s)"] == 'Controller, "Group"\nFinance'
    assert tricky["Email status"] == GUESSED

    # formula-looking cells get a leading apostrophe
    evil = by_company["'=HYPERLINK(\"http://evil.example\",\"click\")"]
    assert evil["Job title(s)"] == "'+cmd|' /C calc'!A0"
    assert evil["Location"] == "'-2+3"
    assert evil["Email status"] == NOT_FOUND and evil["Email"] == ""
    assert "'@SUM Industries" in by_company
    # numbers are not touched
    assert by_company["Acme Corp"]["Score"] == "88"
    assert by_company["Acme Corp"]["Company size"] == "120"


def test_csv_never_writes_internal_keys(tmp_path):
    pkg = make_pkg([make_lead()])
    text = write_csv(pkg, tmp_path / "x.csv").read_text(encoding="utf-8-sig")
    lead_id = pkg.rows[0]["_lead_id"]
    assert lead_id not in text
    assert "_company_key" not in text and "_lead_id" not in text


def test_csv_with_opening_column_and_zero_one_sixty_rows(tmp_path):
    pkg = make_pkg([make_lead()], include_opening=True)
    rows = read_csv(write_csv(pkg, tmp_path / "one.csv"))
    assert rows[0][-1] == "Suggested opening line"
    assert rows[1][-1] == "Saw Acme Corp is hiring."
    assert len(rows) == 2

    empty = read_csv(write_csv(make_pkg([]), tmp_path / "zero.csv"))
    assert empty == [[h for _, h in make_pkg([]).columns]]

    sixty = make_pkg(many_leads(60))
    rows = read_csv(write_csv(sixty, tmp_path / "sixty.csv"))
    assert len(rows) == 61
    assert [r[0] for r in rows[1:]] == [r["company"] for r in sixty.rows]


def test_csv_creates_parent_folders(tmp_path):
    path = write_csv(make_pkg([make_lead()]), tmp_path / "a" / "b" / "c.csv")
    assert path.exists()


# --- XLSX -------------------------------------------------------------------------------

def load(path: Path) -> openpyxl.Workbook:
    return openpyxl.load_workbook(path)


def col_index(pkg: DeliveryPackage, header: str) -> int:
    return [h for _, h in pkg.columns].index(header) + 1


def test_xlsx_structure_header_freeze_filter_widths(tmp_path):
    pkg = make_pkg(tricky_leads(), brand={"brand_color": "#AA3366"})
    wb = load(write_xlsx(pkg, tmp_path / "out.xlsx"))
    assert wb.sheetnames == ["Leads", "About"]
    ws = wb["Leads"]
    header = [c.value for c in ws[1]]
    assert header == [h for _, h in pkg.columns]
    for cell in ws[1]:
        assert cell.font.bold
        assert cell.fill.fill_type == "solid"
        assert cell.fill.fgColor.rgb == "FFAA3366"
    assert ws.freeze_panes == "A2"
    last = openpyxl.utils.get_column_letter(len(pkg.columns))
    assert ws.auto_filter.ref == f"A1:{last}{len(pkg.rows) + 1}"
    assert ws.max_row == len(pkg.rows) + 1
    for i in range(1, len(pkg.columns) + 1):
        width = ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width
        assert 10 <= width <= 60, (i, width)
    # a long column is wider than a short one
    widths = {h: ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width
              for i, (_, h) in enumerate(pkg.columns, start=1)}
    assert widths["Company"] > widths["Score"]


def test_xlsx_links_dates_numbers_and_formula_safety(tmp_path):
    leads = tricky_leads()
    pkg = make_pkg(leads)
    ws = load(write_xlsx(pkg, tmp_path / "out.xlsx"))["Leads"]
    rows = {ws.cell(row=r, column=1).value: r for r in range(2, ws.max_row + 1)}

    acme = rows["Acme Corp"]
    site = ws.cell(row=acme, column=col_index(pkg, "Website"))
    job = ws.cell(row=acme, column=col_index(pkg, "Job link"))
    li = ws.cell(row=acme, column=col_index(pkg, "LinkedIn URL"))
    assert site.hyperlink.target == "https://acme.com" and site.value == "https://acme.com"
    assert job.hyperlink.target == "https://jobs.acme.com/123"
    assert li.hyperlink.target == "https://www.linkedin.com/in/jane-doe"
    assert site.font.underline == "single"

    posted = ws.cell(row=acme, column=col_index(pkg, "Date posted"))
    assert isinstance(posted.value, (date, datetime)) and posted.is_date
    assert (posted.value.date() if isinstance(posted.value, datetime) else posted.value) == date(2026, 9, 22)
    score = ws.cell(row=acme, column=col_index(pkg, "Score"))
    assert score.value == 88 and isinstance(score.value, int)
    assert ws.cell(row=acme, column=col_index(pkg, "Company size")).value == 120

    # formula-looking text stays text, exactly as delivered
    evil_row = next(r for name, r in rows.items() if str(name).startswith("=HYPERLINK"))
    evil = ws.cell(row=evil_row, column=1)
    assert evil.value == '=HYPERLINK("http://evil.example","click")'
    assert evil.data_type == "s"
    assert ws.cell(row=evil_row, column=col_index(pkg, "Job title(s)")).value == "+cmd|' /C calc'!A0"
    for row in ws.iter_rows():
        for cell in row:
            assert cell.data_type != "f", cell.coordinate

    # missing fields: no link for javascript: / empty values, bare domains are linked with https://
    at = rows["@SUM Industries"]
    assert ws.cell(row=at, column=col_index(pkg, "Website")).hyperlink is None
    assert ws.cell(row=at, column=col_index(pkg, "Job link")).value is None
    assert ws.cell(row=at, column=col_index(pkg, "Date posted")).value is None
    assert ws.cell(row=at, column=col_index(pkg, "Company size")).value is None
    li2 = ws.cell(row=at, column=col_index(pkg, "LinkedIn URL"))
    assert li2.hyperlink.target == "https://linkedin.com/in/lee-wong"
    # control characters Excel cannot store are dropped instead of crashing
    assert ws.cell(row=at, column=col_index(pkg, "Job title(s)")).value == "Accountant (bell)"

    # unicode survives
    assert "Société Générale Ünïcødé 株式会社 🚀" in rows


def about_sections(ws: Any) -> Dict[str, List[List[Any]]]:
    """Section title -> rows (after the header row) until the next blank row."""
    out: Dict[str, List[List[Any]]] = {}
    current: Optional[str] = None
    skip_header = False
    for row in ws.iter_rows(values_only=True):
        a, b = (row + (None, None))[:2]
        if a is None and b is None:
            current = None
            continue
        if current is None and b is None and a in ("Leads by signal", "Leads by email status",
                                                   "Email status legend", "Columns"):
            current, skip_header = a, True
            out[current] = []
            continue
        if current is not None:
            if skip_header:
                skip_header = False
                continue
            out[current].append([a, b])
    return out


def test_xlsx_about_sheet(tmp_path):
    pkg = make_pkg(tricky_leads(), notes=["freshness: jobs posted in the last 7 days"],
                   brand={"brand_name": "Acme Signals", "sender_name": "Sam Seller",
                          "sender_email": "sam@seller.example", "website": "https://seller.example",
                          "footer": "Confidential - for Acme Staffing only"})
    ws = load(write_xlsx(pkg, tmp_path / "out.xlsx"))["About"]
    values = {r[0]: r[1] for r in ws.iter_rows(values_only=True) if r[0]}
    assert ws["A1"].value == "Acme Signals"
    assert values["Prepared for"] == "Acme Staffing Ltd"
    assert values["Period"] == "2026-09-17 to 2026-09-24"
    assert values["Leads in this file"] == 5
    assert values["Hot leads"] == pkg.hot == 2
    assert values["Companies"] == 5
    assert values["Notes"] == "freshness: jobs posted in the last 7 days"
    assert "Sam Seller" in values["Prepared by"] and "sam@seller.example" in values["Prepared by"]

    sections = about_sections(ws)
    assert dict(map(tuple, sections["Leads by signal"])) == pkg.counts_by_signal
    assert dict(map(tuple, sections["Leads by email status"])) == pkg.counts_by_email_status
    assert pkg.counts_by_email_status == {VERIFIED: 1, RISKY: 1, GUESSED: 1, NOT_FOUND: 2}
    legend = sections["Email status legend"]
    assert [r[0] for r in legend] == list(EMAIL_LABELS)
    assert {r[0]: r[1] for r in legend} == EMAIL_STATUS_MEANINGS
    assert "Urgency" in {r[0] for r in sections["Columns"]}
    assert "Suggested opening line" not in {r[0] for r in sections["Columns"]}


def test_email_status_meanings_cover_exactly_the_labels():
    assert list(EMAIL_STATUS_MEANINGS) == list(EMAIL_LABELS)
    assert "guess" in EMAIL_STATUS_MEANINGS[VERIFIED].lower()      # "Never a guessed address"
    assert "pattern" in EMAIL_STATUS_MEANINGS[GUESSED]


def test_xlsx_zero_one_and_sixty_rows(tmp_path):
    wb = load(write_xlsx(make_pkg([]), tmp_path / "zero.xlsx"))
    ws = wb["Leads"]
    assert ws.max_row == 1
    last = openpyxl.utils.get_column_letter(len(make_pkg([]).columns))
    assert ws.auto_filter.ref == f"A1:{last}1"
    about = {r[0]: r[1] for r in wb["About"].iter_rows(values_only=True) if r[0]}
    assert about["Leads in this file"] == 0
    assert about["(none)"] == 0

    one = make_pkg([make_lead()], include_opening=True)
    ws = load(write_xlsx(one, tmp_path / "one.xlsx"))["Leads"]
    assert ws.max_row == 2
    assert ws.cell(row=1, column=len(one.columns)).value == "Suggested opening line"
    assert ws.cell(row=2, column=len(one.columns)).value == "Saw Acme Corp is hiring."

    sixty = make_pkg(many_leads(60))
    ws = load(write_xlsx(sixty, tmp_path / "sixty.xlsx"))["Leads"]
    assert ws.max_row == 61
    assert ws.auto_filter.ref.endswith("61")


def test_xlsx_default_and_invalid_brand_colour(tmp_path):
    for brand in ({}, {"brand_color": "blue"}, {"brand_color": None}):
        ws = load(write_xlsx(make_pkg([make_lead()], brand=brand), tmp_path / "b.xlsx"))["Leads"]
        assert ws["A1"].fill.fgColor.rgb == "FF1F4E79"
        assert ws["A1"].font.color.rgb == "FFFFFFFF"
    # a light brand colour gets dark header text
    ws = load(write_xlsx(make_pkg([make_lead()], brand={"brand_color": "#FFE08A"}), tmp_path / "c.xlsx"))["Leads"]
    assert ws["A1"].font.color.rgb == "FF1C2430"


# --- HTML -------------------------------------------------------------------------------

def kpis(page: str) -> Dict[str, int]:
    return {label: int(v) for v, label in re.findall(
        r'<span class="value">(\d+)</span><span class="label">([^<]+)</span>', page)}


def top_table_companies(page: str) -> List[str]:
    body = page.split('<table class="leads">', 1)[1].split("</table>", 1)[0]
    return [re.sub(r"<[^>]+>", "", m) for m in re.findall(r"<tr><td>(.*?)</td>", body)]


def test_html_branding_kpis_signals_and_footer(tmp_path):
    leads = tricky_leads() + [make_lead("Acme Corp", "acme.com", title="Staff Accountant", tier="normal",
                                        score=60, contact=None)]
    leads.append(make_lead("Fundr", "fundr.example", title="Raised $20M Series B", signal_type="funding",
                           tier="normal", score=55))
    brand = {"brand_name": "Acme Signals", "brand_color": "#0B6E4F", "sender_name": "Sam Seller",
             "sender_email": "sam@seller.example", "website": "https://seller.example",
             "footer": "Confidential - prepared for Acme Staffing"}
    pkg = make_pkg(leads, brand=brand, notes=["Jobs posted in the last 7 days"])
    path = write_html(pkg, tmp_path / "r.html")
    page = path.read_text(encoding="utf-8")
    assert page.startswith("<!DOCTYPE html>")
    assert "<title>Acme Signals - Acme Staffing Ltd - 2026-09-24</title>" in page
    assert "--brand:#0b6e4f" in page
    assert '<p class="eyebrow">Acme Signals</p>' in page
    assert "Prepared for Acme Staffing Ltd" in page
    assert "17 Sep 2026 &ndash; 24 Sep 2026" in page
    assert kpis(page) == {"Leads delivered": 7, "Hot leads": 2, "Verified emails": 2, "Companies": 6}
    # counts by signal type
    assert re.search(r"<td>Hiring</td><td class=\"num\">6</td>", page)
    assert re.search(r"<td>Funding</td><td class=\"num\">1</td>", page)
    # footer
    assert "Prepared by Sam Seller" in page
    assert '<a href="mailto:sam@seller.example">sam@seller.example</a>' in page
    assert '<a href="https://seller.example" rel="noopener noreferrer">seller.example</a>' in page
    assert "Confidential - prepared for Acme Staffing" in page
    assert "All 7 leads, hottest first" in page and "Showing" not in page
    assert "<li>Jobs posted in the last 7 days</li>" in page
    # print CSS, self-contained
    assert "@media print" in page
    assert "<link" not in page and "<script" not in page and "src=" not in page
    assert "@import" not in page and "url(" not in page


def test_html_defaults_when_brand_missing(tmp_path):
    page = render_html(make_pkg([make_lead()], brand={}))
    assert '<p class="eyebrow">Hiring Signal Report</p>' in page
    assert "--brand:#1f4e79" in page
    assert "<img" not in page
    page = render_html(make_pkg([make_lead()], brand={"brand_color": "not-a-colour", "brand_name": "  "}))
    assert "--brand:#1f4e79" in page and "Hiring Signal Report" in page


def test_html_logo_only_when_safe(tmp_path):
    page = render_html(make_pkg([make_lead()], brand={"brand_name": "Acme Signals",
                                                      "logo_url": "https://cdn.example.com/logo.png"}))
    assert '<img src="https://cdn.example.com/logo.png" alt="Acme Signals logo">' in page
    page = render_html(make_pkg([make_lead()], brand={"logo_url": "javascript:alert(1)"}))
    assert "<img" not in page and "javascript:" not in page
    page = render_html(make_pkg([make_lead()], brand={"logo_url": 'https://x.example/l.png" onerror="x'}))
    assert 'onerror="x' not in page


def test_html_escapes_every_value():
    lead = make_lead("<script>alert(1)</script>", "xss.example", title='<img src=x onerror="bad()">',
                     location="Austin & <b>TX</b>")
    pkg = make_pkg([lead], brand={"brand_name": "<b>Brand</b>", "footer": "<i>f</i>",
                                  "sender_name": "A & B"}, notes=["<script>n</script>"],
                   display="<em>Client</em>")
    page = render_html(pkg)
    assert "<script" not in page
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in page
    assert "&lt;img src=x onerror=&quot;bad()&quot;&gt;" in page
    assert "Austin &amp; &lt;b&gt;TX&lt;/b&gt;" in page
    assert "<b>" not in page and "<em>" not in page and "<i>" not in page
    assert "&lt;em&gt;Client&lt;/em&gt;" in page
    assert "Prepared by A &amp; B" in page


def test_html_top_table_is_hot_first_in_row_order():
    leads = many_leads(15)                 # rows 0..4 hot, then normal (as run.py sorts them)
    pkg = make_pkg(leads)
    page = render_html(pkg)
    assert top_table_companies(page) == [r["company"] for r in pkg.rows[:10]]
    assert "Top 10 hottest leads" in page
    assert "Showing 10 of 15 leads" in page
    # a hot row listed after normal ones is still shown first (stable within each group)
    pkg.rows.insert(0, pkg.rows.pop(7))    # a normal row first
    pkg.rows.append(dict(pkg.rows[1], company="Late Hot Co", urgency="hot"))
    names = top_table_companies(render_html(pkg))
    hot = [r["company"] for r in pkg.rows if r["urgency"] == "hot"]
    assert names[:len(hot)] == hot
    assert names[len(hot)] == pkg.rows[0]["company"]
    # custom top
    assert len(top_table_companies(render_html(make_pkg(leads), top=3))) == 3


def test_html_table_cells_and_links():
    page = render_html(make_pkg([make_lead()]))
    assert '<a href="https://acme.com" rel="noopener noreferrer">Acme Corp</a>' in page
    assert '<a href="https://jobs.acme.com/123" rel="noopener noreferrer">Senior Accountant</a>' in page
    assert "<td>2 days ago</td>" in page
    assert "<td>Austin, TX</td>" in page
    assert '<span class="badge u-hot">hot</span>' in page
    assert 'Jane Doe<span class="sub">CFO</span>' in page
    assert '<span class="badge st-verified">verified</span>' in page
    assert "Your lead" in page
    # legend under the table lists the four statuses
    legend = page.split('<ul class="legend">', 1)[1].split("</ul>", 1)[0]
    assert re.findall(r'class="badge st-[a-z]+">([^<]+)</span>', legend) == list(EMAIL_LABELS)


def test_html_zero_rows(tmp_path):
    page = write_html(make_pkg([]), tmp_path / "zero.html").read_text(encoding="utf-8")
    assert "No new leads matched your criteria this period." in page
    assert kpis(page) == {"Leads delivered": 0, "Hot leads": 0, "Verified emails": 0, "Companies": 0}
    assert '<table class="leads">' not in page
    assert "Leads by signal type" not in page


def test_html_light_brand_colour_uses_dark_text():
    page = render_html(make_pkg([make_lead()], brand={"brand_color": "#FFE08A"}))
    assert "--brand-ink:#1c2430" in page
    assert "--brand-text:#1c2430" in page
    page = render_html(make_pkg([make_lead()], brand={"brand_color": "#1F4E79"}))
    assert "--brand-ink:#ffffff" in page and "--brand-text:#1f4e79" in page


# --- helpers + write_all ----------------------------------------------------------------

def test_link_target():
    assert link_target("https://acme.com/x?y=1") == "https://acme.com/x?y=1"
    assert link_target("http://acme.com") == "http://acme.com"
    assert link_target("acme.com") == "https://acme.com"
    assert link_target("www.linkedin.com/in/jane") == "https://www.linkedin.com/in/jane"
    for bad in ("", None, "javascript:alert(1)", "mailto:a@b.com", "not a url", "https://", "ftp://x.com",
                "data:text/html,hi"):
        assert link_target(bad) == "", bad


@pytest.mark.parametrize("bad", ["https://acme.com]", "https://jobs.acme.com]/123", "https://[acme.com]",
                                 "http://[::1", "https://acme.com[/x", "acme.com]"])
def test_link_target_never_raises_on_a_malformed_url(bad):
    # urlparse raises ValueError on an unbalanced / non-IP bracketed host; such a value is shown as text
    assert link_target(bad) == ""


def test_malformed_urls_do_not_crash_any_format(tmp_path):
    """A stray bracket in a website / job URL / logo URL (sources let these through) must not stop the
    delivery after the pipeline has run: the value is written as plain text without a link."""
    leads = [make_lead(website="https://acme.com]"),
             make_lead("Beta LLC", "beta.example", url="https://jobs.beta.example]/123"),
             make_lead("Gamma Inc", "gamma.example", website="https://[gamma.example]")]
    pkg = make_pkg(leads, brand={"logo_url": "https://[logo.example]/logo.png"})
    out = write_all(pkg, tmp_path, ["csv", "xlsx", "html"])
    assert set(out) == {"csv", "xlsx", "html"} and all(p.is_file() for p in out.values())
    ws = openpyxl.load_workbook(out["xlsx"])["Leads"]
    header = [c.value for c in ws[1]]
    web, job = header.index("Website") + 1, header.index("Job link") + 1
    assert ws.cell(row=2, column=web).value == "https://acme.com]" and ws.cell(row=2, column=web).hyperlink is None
    assert ws.cell(row=3, column=job).value == "https://jobs.beta.example]/123"
    assert ws.cell(row=3, column=job).hyperlink is None
    page = out["html"].read_text(encoding="utf-8")
    assert "<img" not in page and 'href="https://acme.com]"' not in page and "Acme Corp" in page


def test_write_all_leaves_no_partial_files_when_a_format_fails(tmp_path, monkeypatch):
    """A writer failing half-way must not leave a complete-looking (but never recorded) client file."""
    def boom(pkg, path):
        Path(path).write_text("half a workbook", encoding="utf-8")
        raise RuntimeError("disk full")

    monkeypatch.setitem(formats.WRITERS, "xlsx", boom)
    folder = tmp_path / "2026-09-24"
    with pytest.raises(RuntimeError, match="disk full"):
        write_all(make_pkg([make_lead()]), folder, ["csv", "xlsx", "html"])
    assert list(folder.iterdir()) == []


def test_brand_settings_merges_defaults():
    pkg = make_pkg([], brand={"brand_name": "X", "sender_name": " Sam ", "extra": 5})
    b = brand_settings(pkg)
    assert b["brand_name"] == "X" and b["sender_name"] == "Sam" and b["brand_color"] == "#1f4e79"
    assert b["extra"] == "5" and b["logo_url"] == ""


def test_write_all_names_and_order(tmp_path):
    pkg = make_pkg([make_lead()])
    out = write_all(pkg, tmp_path / "deliveries" / "acme" / "2026-09-24", ["xlsx", "CSV", "html", "csv"])
    assert list(out) == ["xlsx", "csv", "html"]
    assert {k: p.name for k, p in out.items()} == {
        "xlsx": "acme-hiring-signals-2026-09-24.xlsx",
        "csv": "acme-hiring-signals-2026-09-24.csv",
        "html": "acme-hiring-signals-2026-09-24.html",
    }
    assert all(p.exists() and p.stat().st_size > 0 for p in out.values())
    preview = write_all(pkg, tmp_path, "csv", preview=True)
    assert preview["csv"].name == "acme-hiring-signals-2026-09-24-PREVIEW.csv"
    assert write_all(pkg, tmp_path, []) == {}


def test_write_all_rejects_unknown_format(tmp_path):
    with pytest.raises(ValueError, match=r"'pdf'.*csv, xlsx, html"):
        write_all(make_pkg([]), tmp_path, ["csv", "pdf"])
    assert not list(tmp_path.iterdir())     # nothing written when the list is invalid


def test_file_name_is_file_safe():
    pkg = make_pkg([], client="../Acme Staffing/ltd")
    assert file_name(pkg, "csv") == "Acme-Staffing-ltd-hiring-signals-2026-09-24.csv"
    assert file_name(make_pkg([], client=""), ".html", preview=True) == "client-hiring-signals-2026-09-24-PREVIEW.html"


# --- Google Sheets (fake gspread) -------------------------------------------------------

class WorksheetNotFound(Exception):
    pass


class SpreadsheetNotFound(Exception):
    pass


class FakeWorksheet:
    def __init__(self, title: str, rows: int = 1000, cols: int = 26, ws_id: Optional[int] = None):
        self.title, self.row_count, self.col_count = title, rows, cols
        self.values: List[List[Any]] = []
        self.calls: List[tuple] = []
        if ws_id is not None:
            self.id = ws_id

    def clear(self):
        self.calls.append(("clear",))
        self.values = []

    def resize(self, rows=None, cols=None):
        self.calls.append(("resize", rows, cols))
        self.row_count, self.col_count = rows or self.row_count, cols or self.col_count

    def update(self, range_name=None, values=None, value_input_option=None, **kw):
        """Like the Sheets API: writes the block at A1 over the cells it covers, keeps the rest."""
        self.calls.append(("update", range_name, value_input_option))
        if len(values) > self.row_count or max(len(r) for r in values) > self.col_count:
            raise AssertionError("exceeds grid limits")
        assert range_name == "A1"
        grid = [list(r) for r in self.values]
        for i, row in enumerate(values):
            if i >= len(grid):
                grid.append([])
            line = grid[i]
            line.extend([""] * (len(row) - len(line)))
            line[:len(row)] = list(row)
        self.values = grid

    def get_all_values(self):
        self.calls.append(("read",))
        return [list(r) for r in self.values]

    def shown(self):
        """What a person sees: the values without trailing empty cells / rows."""
        rows = [list(r) for r in self.values]
        for r in rows:
            while r and r[-1] in ("", None):
                r.pop()
        while rows and not rows[-1]:
            rows.pop()
        return rows

    def freeze(self, rows=None, cols=None):
        self.calls.append(("freeze", rows))


class FakeBook:
    def __init__(self, key: str, sheets: List[FakeWorksheet]):
        self.id = key
        self.url = f"https://docs.google.com/spreadsheets/d/{key}"
        self.sheets = {ws.title: ws for ws in sheets}
        self.added: List[tuple] = []

    def worksheet(self, title):
        if title not in self.sheets:
            raise WorksheetNotFound(title)
        return self.sheets[title]

    def add_worksheet(self, title, rows, cols):
        self.added.append((title, rows, cols))
        ws = self.sheets[title] = FakeWorksheet(title, rows, cols)
        return ws


class FakeClient:
    def __init__(self, book: FakeBook):
        self.book = book
        self.opened: List[tuple] = []

    def open_by_key(self, key):
        self.opened.append(("key", key))
        if key != self.book.id:
            raise SpreadsheetNotFound(key)
        return self.book

    def open_by_url(self, url):
        self.opened.append(("url", url))
        if self.book.id not in url:
            raise SpreadsheetNotFound(url)
        return self.book


def fake_gspread(client: FakeClient) -> types.ModuleType:
    mod = types.ModuleType("gspread")
    exc = types.ModuleType("gspread.exceptions")
    exc.WorksheetNotFound = WorksheetNotFound
    exc.SpreadsheetNotFound = SpreadsheetNotFound
    mod.exceptions = exc
    mod.auth_calls = []

    def service_account(filename=None, **kw):
        mod.auth_calls.append(("file", filename))
        return client

    def service_account_from_dict(info, **kw):
        mod.auth_calls.append(("dict", info))
        return client

    mod.service_account = service_account
    mod.service_account_from_dict = service_account_from_dict
    return mod


@pytest.fixture
def sheets(monkeypatch, tmp_path):
    key_file = tmp_path / "sa.json"
    key_file.write_text('{"type": "service_account"}', encoding="utf-8")
    book = FakeBook("1AbCdEfG", [FakeWorksheet("2026-09-24", rows=50, cols=10, ws_id=777)])
    client = FakeClient(book)
    mod = fake_gspread(client)
    monkeypatch.setitem(sys.modules, "gspread", mod)
    return SimpleNamespace(book=book, client=client, mod=mod, key_file=str(key_file))


def test_push_google_sheet_replaces_worksheet(make_ctx, sheets):
    """A fixed tab name (no {date}) is a rolling tab: each delivery replaces it."""
    ctx = make_ctx(mode="delivery")
    ws = sheets.book.sheets["Leads"] = FakeWorksheet("Leads", rows=50, cols=10, ws_id=777)
    ws.values = [["old"], ["junk"]]
    pkg = make_pkg(many_leads(60), include_opening=True)
    url = push_google_sheet(pkg, {"spreadsheet_id": "1AbCdEfG", "worksheet": "Leads",
                                  "service_account_file": sheets.key_file}, ctx)
    assert url == "https://docs.google.com/spreadsheets/d/1AbCdEfG#gid=777"
    assert sheets.mod.auth_calls == [("file", sheets.key_file)]
    assert sheets.client.opened == [("key", "1AbCdEfG")]
    assert sheets.book.added == []
    # 61 rows x 19 columns > the 50 x 10 grid: grown before writing, RAW values; never cleared first
    assert [c[0] for c in ws.calls] == ["read", "resize", "update", "freeze"]
    assert ws.calls[1] == ("resize", 61, 19) and ws.calls[2] == ("update", "A1", "RAW")
    header = [h for _, h in pkg.columns]
    assert ws.values[0] == header
    assert ws.values[1:] == pkg.public_rows()
    first = dict(zip(header, ws.values[1]))
    assert first["Company"] == pkg.rows[0]["company"] and first["Score"] == pkg.rows[0]["score"]
    assert isinstance(first["Score"], int)


def test_push_google_sheet_creates_missing_worksheet(make_ctx, sheets):
    ctx = make_ctx(mode="delivery", env={"GOOGLE_APPLICATION_CREDENTIALS": sheets.key_file})
    pkg = make_pkg(tricky_leads())
    url = push_google_sheet(pkg, {"spreadsheet_url": "https://docs.google.com/spreadsheets/d/1AbCdEfG/edit",
                                  "worksheet": "{client} {date}"}, ctx)
    assert sheets.book.added == [("acme 2026-09-24", 100, 26)]
    ws = sheets.book.sheets["acme 2026-09-24"]
    assert url == "https://docs.google.com/spreadsheets/d/1AbCdEfG"     # new worksheet has no id in the fake
    assert sheets.client.opened == [("url", "https://docs.google.com/spreadsheets/d/1AbCdEfG/edit")]
    assert sheets.mod.auth_calls == [("file", sheets.key_file)]
    # RAW: formula-looking text is sent as-is (Sheets will not evaluate it), no apostrophe added
    assert any(row[0] == '=HYPERLINK("http://evil.example","click")' for row in ws.values)
    assert all(v is not None for row in ws.values for v in row)


def test_push_google_sheet_default_worksheet_and_json_credentials(make_ctx, sheets):
    ctx = make_ctx(mode="delivery", env={"GOOGLE_SERVICE_ACCOUNT_JSON": '{"type": "service_account"}'})
    push_google_sheet(make_pkg([]), {"spreadsheet_id": "1AbCdEfG"}, ctx)
    assert sheets.mod.auth_calls == [("dict", {"type": "service_account"})]
    assert sheets.book.sheets["2026-09-24"].values == [[h for _, h in make_pkg([]).columns]]


def test_push_google_sheet_replacing_a_bigger_tab_blanks_the_leftover_cells(make_ctx, sheets):
    ctx = make_ctx(mode="delivery", env={"GOOGLE_SERVICE_ACCOUNT_JSON": '{"type": "service_account"}'})
    ws = sheets.book.sheets["Leads"] = FakeWorksheet("Leads", rows=100, cols=30)
    ws.values = [[f"old {r}-{c}" for c in range(25)] for r in range(40)]     # last week: 40 rows x 25 cols
    pkg = make_pkg([make_lead()])
    push_google_sheet(pkg, {"spreadsheet_id": "1AbCdEfG", "worksheet": "Leads"}, ctx)
    assert ws.shown() == [[h for _, h in pkg.columns]] + pkg.public_rows()
    assert "clear" not in [c[0] for c in ws.calls]


def test_push_google_sheet_failed_write_keeps_the_old_contents(make_ctx, sheets):
    """The new values are written in one call before anything is removed: when that call fails,
    the client's tab still holds what it held (it is never left empty)."""
    ctx = make_ctx(mode="delivery", env={"GOOGLE_SERVICE_ACCOUNT_JSON": '{"type": "service_account"}'})
    ws = sheets.book.sheets["Leads"] = FakeWorksheet("Leads", rows=50, cols=30)
    before = [["Company"], ["Lead delivered last week"]]
    ws.values = [list(r) for r in before]

    def quota(**kw):
        raise RuntimeError("quota exceeded")

    ws.update = quota
    with pytest.raises(SheetPushError, match="quota exceeded"):
        push_google_sheet(make_pkg([make_lead()]), {"spreadsheet_id": "1AbCdEfG", "worksheet": "Leads"}, ctx)
    assert ws.values == before


def test_push_google_sheet_never_overwrites_an_earlier_delivery_tab(make_ctx, sheets):
    """A {date} tab is one delivery. A second delivery on the same day (its files go to <folder>-2)
    or another client using the same spreadsheet must not wipe it: the rows go to '<date>-2'."""
    ctx = make_ctx(mode="delivery", env={"GOOGLE_SERVICE_ACCOUNT_JSON": '{"type": "service_account"}'})
    cfg = {"spreadsheet_id": "1AbCdEfG"}                        # default worksheet: "{date}"
    first = make_pkg([make_lead("First Co", "first.example"), make_lead("Second Co", "second.example")])
    push_google_sheet(first, cfg, ctx)
    tab = sheets.book.sheets["2026-09-24"]
    assert tab.shown() == [[h for _, h in first.columns]] + first.public_rows()

    again = make_pkg([make_lead("Third Co", "third.example")])
    push_google_sheet(again, cfg, ctx)
    assert tab.shown() == [[h for _, h in first.columns]] + first.public_rows()     # untouched
    assert "update" not in [c[0] for c in tab.calls[-2:]]
    assert sheets.book.sheets["2026-09-24-2"].shown() == [[h for _, h in again.columns]] + again.public_rows()

    third = make_pkg([make_lead("Fourth Co", "fourth.example")], client="other")
    push_google_sheet(third, cfg, ctx)
    assert sheets.book.sheets["2026-09-24-3"].shown()[1][0] == "Fourth Co"
    assert [t for t, _, _ in sheets.book.added] == ["2026-09-24-2", "2026-09-24-3"]
    # an existing but empty tab for the date is simply used
    assert formats.worksheet_title("{client} {date}", third) == "other 2026-09-24"
    sheets.book.sheets["other 2026-09-24"] = empty = FakeWorksheet("other 2026-09-24")
    push_google_sheet(third, {"spreadsheet_id": "1AbCdEfG", "worksheet": "{client} {date}"}, ctx)
    assert empty.shown()[1][0] == "Fourth Co" and "other 2026-09-24-2" not in sheets.book.sheets


def test_push_google_sheet_errors(make_ctx, sheets, tmp_path):
    ctx = make_ctx(mode="delivery")
    pkg = make_pkg([make_lead()])
    with pytest.raises(ValueError, match="spreadsheet_id"):
        push_google_sheet(pkg, {}, ctx)
    with pytest.raises(SheetPushError, match="GOOGLE_APPLICATION_CREDENTIALS"):
        push_google_sheet(pkg, {"spreadsheet_id": "1AbCdEfG"}, ctx)
    with pytest.raises(SheetPushError, match="key file not found"):
        push_google_sheet(pkg, {"spreadsheet_id": "1AbCdEfG",
                                "service_account_file": str(tmp_path / "missing.json")}, ctx)
    with pytest.raises(SheetPushError, match="not found - check the id"):
        push_google_sheet(pkg, {"spreadsheet_id": "nope", "service_account_file": sheets.key_file}, ctx)
    with pytest.raises(SheetPushError, match="not valid JSON"):
        push_google_sheet(pkg, {"spreadsheet_id": "1AbCdEfG", "service_account_json": "{nope"}, ctx)


def test_push_google_sheet_dry_run_does_nothing(make_ctx, sheets):
    ctx = make_ctx(mode="delivery", dry_run=True)
    assert push_google_sheet(make_pkg([make_lead()]), {"spreadsheet_id": "1AbCdEfG",
                                                       "service_account_file": sheets.key_file}, ctx) == ""
    assert sheets.mod.auth_calls == [] and sheets.client.opened == []


def test_push_google_sheet_without_gspread(make_ctx, monkeypatch, tmp_path):
    monkeypatch.setitem(sys.modules, "gspread", None)   # import gspread -> ImportError
    ctx = make_ctx(mode="delivery")
    with pytest.raises(RuntimeError, match="pip install gspread"):
        push_google_sheet(make_pkg([make_lead()]), {"spreadsheet_id": "1AbCdEfG"}, ctx)


def test_worksheet_title_limits():
    pkg = make_pkg([])
    assert formats.worksheet_title("", pkg) == "2026-09-24"
    assert formats.worksheet_title("  Week   {date} ", pkg) == "Week 2026-09-24"
    assert len(formats.worksheet_title("x" * 300, pkg)) == 100
