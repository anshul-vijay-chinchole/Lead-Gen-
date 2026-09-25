"""End-to-end: the lead-delivery business through the real CLI (``main([...])``).

What is checked, with only the shipped sample data (``clients/demo-client.yaml``
on ``playbooks/demo-delivery.yaml``), a temporary database, temporary output
folders, an empty env file and a ``FakeHttp`` in place of the network:

* ``leadgen deliver --client demo-client`` writes valid client files: the CSV
  header is exactly ``rows.columns(include_opening=True)``, the Excel workbook
  opens (sheets "Leads" + "About"), the HTML page names the client.
* Email honesty: every status is one of ``rows.EMAIL_LABELS``, and an address
  built from a name pattern is never labelled "verified" - not even when a paid
  checker (MillionVerifier, faked) says it is deliverable.
* Freshness: every job recorded as delivered was posted within the client's
  ``freshness_days``; stale, undated and re-posted ads never are.
* The ledger: a second delivery contains no company, job or contact from the
  first; a third finds nothing new (exit 1); a dry run records nothing.
* ``replies`` / ``serve`` / ``followups`` refuse to run in delivery mode
  (exit 2, nothing opened), while the outbound demo (``demo-offline``) still
  runs end to end with email sequences written and replies handled.
* The docs do not drift: every ``leadgen ...`` command shown in README.md,
  PLAN.md and docs/ARCHITECTURE.md exists with the flags it shows, every
  command is in the README's command table, the README's ``#section`` links
  resolve, and ``.env.example`` lists every variable the code and the shipped
  playbooks read (with no secret values filled in).

The sample jobs use relative dates ("2 days ago") and the CLI runs "as of
today", so the expectations hold whatever day the tests run.
"""
from __future__ import annotations

import csv
import json
import re
import sqlite3
from datetime import date
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

import openpyxl
import pytest
import yaml

from leadgen import cli, registry
from leadgen.cli import build_parser, main
from leadgen.delivery import rows as R
from leadgen.delivery.client import load_client
from leadgen.delivery.ledger import Ledger
from leadgen.store import Store
from tests.fakes import FakeHttp

REPO = Path(__file__).resolve().parent.parent
CLIENT = "demo-client"
DELIVERY_DEMO = "playbooks/demo-delivery.yaml"
LIVE_DELIVERY = "playbooks/recruitment-delivery.yaml"
OUTBOUND_DEMO = "playbooks/demo-offline.yaml"
REPLIES = "examples/data/demo_replies.csv"
DEMO_JOBS = REPO / "examples" / "data" / "demo_jobs.csv"
DEMO_CONTACTS = REPO / "examples" / "data" / "demo_delivery_contacts.csv"
MV_URL = "https://api.millionverifier.com/api/v3/"
MV_KEY = "mv-e2e-secret-key"

# Every variable an adapter or playbook reads: cleared so a developer's real keys never leak in.
_ENV_VARS = ("THEIRSTACK_API_KEY", "ADZUNA_APP_ID", "ADZUNA_APP_KEY", "APOLLO_API_KEY", "APIFY_TOKEN",
             "HUNTER_API_KEY", "MILLIONVERIFIER_API_KEY", "ZEROBOUNCE_API_KEY", "NEVERBOUNCE_API_KEY",
             "OPENAI_API_KEY", "ANTHROPIC_API_KEY", "INSTANTLY_API_KEY", "SMARTLEAD_API_KEY",
             "SLACK_WEBHOOK_URL", "LEADGEN_WEBHOOK_URL", "LEADGEN_EXPORT_WEBHOOK_URL",
             "GOOGLE_APPLICATION_CREDENTIALS", "GOOGLE_SERVICE_ACCOUNT_JSON", "LEADGEN_WEBHOOK_TOKEN",
             "SENDER_NAME", "SENDER_TITLE", "SENDER_COMPANY", "SENDER_WEBSITE", "SENDER_EMAIL", "BOOKING_LINK",
             "INSTANTLY_CAMPAIGN_ID", "SMARTLEAD_CAMPAIGN_ID", "LEADGEN_PLUGINS")


# --- fixtures + helpers -----------------------------------------------------------------------------

@pytest.fixture(autouse=True)
def _repo_root(monkeypatch):
    """Run from the repository root (as the README says) with no real keys in the environment."""
    monkeypatch.chdir(REPO)
    for k in _ENV_VARS:
        monkeypatch.delenv(k, raising=False)


@pytest.fixture
def http(monkeypatch) -> FakeHttp:
    """The only HTTP client any command gets. No routes: an unexpected request fails the test."""
    fake = FakeHttp()
    monkeypatch.setattr(cli, "make_http", lambda: fake)
    return fake


@pytest.fixture
def ws(tmp_path, http) -> Dict[str, Any]:
    env = tmp_path / "empty.env"
    env.write_text("", encoding="utf-8")
    db = tmp_path / "ledger.db"
    return {"tmp": tmp_path, "db": db, "http": http, "g": ["--db", str(db), "--env-file", str(env)]}


def deliver(ws: Dict[str, Any], out_name: str, *extra: str, client: str = CLIENT) -> Tuple[int, Path]:
    """``leadgen deliver --client <client> --out <tmp>/<out_name> ...`` -> (exit code, folder)."""
    folder = ws["tmp"] / out_name
    code = main(["deliver", "--client", client, "--out", str(folder), *ws["g"], *extra])
    return code, folder


def client_files(folder: Path, preview: bool = False) -> Dict[str, Path]:
    """The client-facing files in a delivery folder, by extension (PREVIEW files or real ones)."""
    out: Dict[str, Path] = {}
    for p in sorted(folder.glob("*-hiring-signals-*.*")):
        if ("-PREVIEW." in p.name) == preview:
            out[p.suffix.lstrip(".")] = p
    return out


def read_rows(csv_path: Path) -> Tuple[List[str], List[Dict[str, str]]]:
    """(header, rows) of a delivery CSV (UTF-8 with BOM)."""
    with open(csv_path, newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        header = next(reader)
        return header, [dict(zip(header, r)) for r in reader]


def days_ago(text: str) -> Optional[int]:
    """'today' -> 0, '1 day ago' -> 1, '24 days ago' -> 24, '' -> None (the sample data's date style)."""
    t = (text or "").strip().lower()
    if not t:
        return None
    if t == "today":
        return 0
    m = re.fullmatch(r"(\d+) days? ago", t)
    assert m, f"unexpected sample date {text!r}"
    return int(m.group(1))


def sample_jobs() -> Dict[str, Dict[str, Any]]:
    """Job ID -> {company, title, age (days, None = undated)} from the shipped sample jobs."""
    with open(DEMO_JOBS, newline="", encoding="utf-8") as f:
        return {r["Job ID"]: {"company": r["Company"], "title": r["Job Title"], "age": days_ago(r["Date Posted"])}
                for r in csv.DictReader(f)}


def provider_emails() -> Dict[str, str]:
    """Email -> status for the addresses the sample contact list supplies (not pattern guesses)."""
    with open(DEMO_CONTACTS, newline="", encoding="utf-8") as f:
        return {r["Email"].strip().lower(): r["Email Status"].strip().lower()
                for r in csv.DictReader(f) if r["Email"].strip()}


def ledger_items(db: Path, client: str = CLIENT) -> Dict[str, Dict[str, Set[str]]]:
    """run_id -> kind -> keys recorded in the delivery ledger for this client."""
    out: Dict[str, Dict[str, Set[str]]] = {}
    with sqlite3.connect(str(db)) as conn:
        for kind, key, run_id in conn.execute(
                "SELECT kind, key, run_id FROM deliveries WHERE client=?", (client,)):
            out.setdefault(run_id, {}).setdefault(kind, set()).add(key)
    return out


def contact_identities(rows: List[Dict[str, str]]) -> Set[str]:
    """Every way a delivered person can be recognised: email, LinkedIn URL, name at company."""
    ids: Set[str] = set()
    for r in rows:
        if r["Email"]:
            ids.add("email:" + r["Email"].lower())
        if r["LinkedIn URL"]:
            ids.add("li:" + r["LinkedIn URL"].lower().rstrip("/"))
        if r["Decision-maker"]:
            ids.add(f"name:{r['Decision-maker'].lower()}|{r['Company'].lower()}")
    return ids


# --- 1. the first delivery: valid client files --------------------------------------------------------

def test_first_delivery_writes_valid_client_files(ws, capsys):
    client = load_client(CLIENT)
    code, folder = deliver(ws, "week1")
    out = capsys.readouterr().out
    assert code == 0, out

    # the three client-ready files, named <client>-hiring-signals-<date>.<ext>
    files = client_files(folder)
    today = date.today().isoformat()
    assert set(files) == {"csv", "xlsx", "html"}
    for ext, path in files.items():
        assert path.name == f"{CLIENT}-hiring-signals-{today}.{ext}"
    assert client_files(folder, preview=True) == {}

    # CSV: Excel-friendly BOM, header == the contract's columns (the demo client has opening lines on)
    raw = files["csv"].read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    header, rows = read_rows(files["csv"])
    assert client.opening_line.enabled is True
    assert header == [h for _, h in R.columns(include_opening=True)]
    assert not any(h.startswith("_") for h in header)            # internal keys never reach the client
    assert len(rows) == client.leads_per_week == 10              # the demo has 12 matches: 2 held back
    assert len({r["Company"] for r in rows}) == len(rows)        # one row per company
    for r in rows:
        assert r["Suggested opening line"].startswith(f"Saw {r['Company']} is hiring"), r
        assert r["Urgency"] in client.tiers
        assert r["Signal type"] == "Hiring"
    # hot leads first, then by score
    order = [(0 if r["Urgency"] == "hot" else 1, -int(r["Score"])) for r in rows]
    assert order == sorted(order)

    # Excel: opens, "Leads" has the same header + rows, "About" carries the email legend
    wb = openpyxl.load_workbook(files["xlsx"])
    assert {"Leads", "About"} <= set(wb.sheetnames)
    sheet = list(wb["Leads"].iter_rows(values_only=True))
    assert list(sheet[0]) == header
    assert [row[0] for row in sheet[1:]] == [r["Company"] for r in rows]
    about = " ".join(str(v) for row in wb["About"].iter_rows(values_only=True) for v in row if v is not None)
    assert client.display_name in about
    for label in R.EMAIL_LABELS:
        assert label in about

    # HTML: one self-contained page that names the client
    html = files["html"].read_text(encoding="utf-8")
    assert "Northstar Finance Recruiting" == client.display_name
    assert client.display_name in html and "Hiring Signal Report" in html
    assert "<script" not in html.lower()

    # internal files: QA summary + everything left out, and why
    internal = folder / "_internal"
    qa = json.loads((internal / "qa.json").read_text(encoding="utf-8"))
    assert qa["delivered"] == len(rows) and qa["target"] == 10 and qa["dry_run"] is False
    assert qa["held_back"] == 2
    assert (internal / "qa.txt").is_file()
    with open(internal / "not_delivered.csv", newline="", encoding="utf-8-sig") as f:
        left_out = {r["company"]: r["reason"] for r in csv.DictReader(f)}
    assert "on this client's do-not-list" in left_out["Riverwalk Hospitality Group"]   # client exclusion
    assert "excluded domain" in left_out["Gulf Coast Paper Co"]
    assert "staffing" in left_out["Summit Staffing Group"]                            # a competitor

    # what the owner sees, in order: files, internal files, QA, next step
    positions = [out.index(s) for s in ("Files for the client", "Internal files", "Delivery QA", "Next: send")]
    assert positions == sorted(positions)
    assert ws["http"].calls == []                                  # 100% offline


# --- 2. email honesty -----------------------------------------------------------------------------------

def test_email_labels_are_honest(ws):
    code, folder = deliver(ws, "week1")
    assert code == 0
    _, rows = read_rows(client_files(folder)["csv"])
    supplied = provider_emails()

    labels = [r["Email status"] for r in rows]
    assert set(labels) <= set(R.EMAIL_LABELS)
    assert set(labels) == set(R.EMAIL_LABELS)                      # the demo shows all four
    for r in rows:
        email = r["Email"].lower()
        assert bool(email) == (r["Email status"] != R.NOT_FOUND), r
        if email and email not in supplied:                       # built by the pattern finder
            assert r["Email status"] == R.GUESSED, r
        if r["Email status"] == R.VERIFIED:                       # only the provider can confirm offline
            assert supplied.get(email) == "verified", r

    # cross-check with what the pipeline stored: a guessed contact is never "verified"
    store = Store(str(ws["db"]))
    try:
        guessed = 0
        for r in rows:
            if not r["Email"]:
                continue
            lead = store.find_lead_by_email(r["Email"])
            assert lead is not None and lead.contact is not None, r
            if R.is_guessed(lead.contact):
                guessed += 1
                assert r["Email status"] == R.GUESSED, r
        assert guessed >= 1
    finally:
        store.close()


def _write_mv_client(clients_dir: Path) -> None:
    """A copy of the demo client whose base playbook checks every email with MillionVerifier."""
    data = yaml.safe_load((REPO / "clients" / f"{CLIENT}.yaml").read_text(encoding="utf-8"))
    data["playbook"] = str(REPO / DELIVERY_DEMO)
    data["display_name"] = "Checker Test Recruiting"
    data["overrides"] = {"enrichment": {"verifier": {"type": "millionverifier"}}}
    clients_dir.mkdir(parents=True, exist_ok=True)
    (clients_dir / "mv-client.yaml").write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")


def test_guessed_email_is_never_verified_even_when_a_paid_checker_says_valid(tmp_path, http):
    """MillionVerifier answers 'ok' (deliverable) for EVERY address, pattern guesses included.
    Provider-supplied addresses may then be 'verified'; a guess must stay 'guessed-unverified'."""
    asked: List[str] = []

    def answer(call: Dict[str, Any]) -> Dict[str, Any]:
        email = call["params"]["email"]
        asked.append(email)
        return {"email": email, "quality": "good", "result": "ok", "resultcode": 1, "subresult": "ok",
                "free": False, "role": False, "didyoumean": "", "credits": 9000 - len(asked),
                "executiontime": 1, "error": "", "livemode": True}

    http.add("GET", MV_URL, fn=answer)
    _write_mv_client(tmp_path / "clients")
    env = tmp_path / "keys.env"
    env.write_text(f"MILLIONVERIFIER_API_KEY={MV_KEY}\n", encoding="utf-8")
    folder = tmp_path / "out"
    code = main(["deliver", "--client", "mv-client", "--clients-dir", str(tmp_path / "clients"),
                 "--out", str(folder), "--db", str(tmp_path / "t.db"), "--env-file", str(env)])
    assert code == 0

    # request shape: the documented single-check endpoint, key + email as query parameters
    assert http.calls and asked
    for call in http.calls:
        assert call["method"] == "GET" and call["url"] == MV_URL
        assert call["params"]["api"] == MV_KEY and call["params"]["email"]

    _, rows = read_rows(client_files(folder)["csv"])
    supplied = provider_emails()
    guessed_rows = [r for r in rows if r["Email"] and r["Email"].lower() not in supplied]
    assert guessed_rows, "the sample data should produce at least one pattern-guessed address"
    for r in guessed_rows:
        assert r["Email"] in asked                                 # the checker DID say "ok" for it ...
        assert r["Email status"] == R.GUESSED, r                   # ... and it is still not "verified"
    assert any(r["Email status"] == R.VERIFIED for r in rows)      # supplied + confirmed = verified
    for r in rows:
        if r["Email status"] == R.VERIFIED:
            assert r["Email"].lower() in supplied, r


# --- 3. freshness -----------------------------------------------------------------------------------------

def test_every_delivered_job_is_within_the_freshness_window(ws):
    client = load_client(CLIENT)
    code, folder = deliver(ws, "week1")
    assert code == 0
    _, rows = read_rows(client_files(folder)["csv"])
    today = date.today()

    # the visible columns: a real posting date inside the window, and a matching "Posted" phrase
    assert client.allow_undated is False
    for r in rows:
        assert r["Date posted"], r
        age = (today - date.fromisoformat(r["Date posted"])).days
        assert 0 <= age <= client.freshness_days, r
        assert r["Posted"] == ("posted today" if age == 0 else f"posted {age} day{'s' if age != 1 else ''} ago")

    # every job recorded as delivered (not only the one linked per row) is fresh
    jobs = sample_jobs()
    recorded = set().union(*(kinds.get("job", set()) for kinds in ledger_items(ws["db"]).values()))
    job_ids = {key.rsplit("|", 1)[1] for key in recorded if key.rsplit("|", 1)[1] in jobs}
    assert len(job_ids) >= len(rows)
    for job_id in job_ids:
        age = jobs[job_id]["age"]
        assert age is not None and age <= client.freshness_days, (job_id, jobs[job_id])
    # the stale ones in the sample data never went out: too old, the re-posted ad, the undated job
    for stale in ("LF-187", "LM-088", "OM-45", "OM-41", "GL-7"):
        assert stale not in job_ids
    # and no row lists a role title that only exists as a stale posting at that company
    for r in rows:
        for title in [t.strip() for t in re.sub(r"\s*\(\+\d+ more\)$", "", r["Job title(s)"]).split(";")]:
            ages = [j["age"] for j in jobs.values() if j["company"] == r["Company"] and j["title"] == title]
            assert any(a is not None and a <= client.freshness_days for a in ages), (r["Company"], title)


# --- 4. the ledger: never the same company, job or contact twice -------------------------------------------

def test_next_delivery_never_repeats_a_company_job_or_contact(ws, capsys):
    code1, week1 = deliver(ws, "week1")
    code2, week2 = deliver(ws, "week2")
    assert (code1, code2) == (0, 0)
    _, rows1 = read_rows(client_files(week1)["csv"])
    _, rows2 = read_rows(client_files(week2)["csv"])
    assert len(rows1) == 10
    assert len(rows2) >= 1                                          # the two leads held back in week 1

    # nothing visible repeats: company, website, job link, person
    assert not {r["Company"] for r in rows1} & {r["Company"] for r in rows2}
    assert not ({r["Website"] for r in rows1} - {""}) & ({r["Website"] for r in rows2} - {""})
    assert not ({r["Job link"] for r in rows1} - {""}) & ({r["Job link"] for r in rows2} - {""})
    assert not contact_identities(rows1) & contact_identities(rows2)

    # nothing recorded repeats either: the ledger's company / job / contact keys per delivery run
    runs = ledger_items(ws["db"])
    assert len(runs) == 2
    first, second = (json.loads((week / "_internal" / "qa.json").read_text(encoding="utf-8"))["run_id"]
                     for week in (week1, week2))
    assert set(runs) == {first, second}
    for kind in ("company", "job", "contact"):
        assert runs[first].get(kind) and runs[second].get(kind), kind
        assert not runs[first][kind] & runs[second][kind], kind
    assert len(runs[first]["company"]) == len(rows1) and len(runs[second]["company"]) == len(rows2)

    store = Store(str(ws["db"]))
    try:
        summary = Ledger(store).summary(CLIENT)
    finally:
        store.close()
    assert summary["deliveries"] == 2

    # week 3: nothing new -> exit 1, empty files, and the reasons say why
    capsys.readouterr()
    code3, week3 = deliver(ws, "week3")
    captured = capsys.readouterr()
    assert code3 == 1
    _, rows3 = read_rows(client_files(week3)["csv"])
    assert rows3 == []
    assert "already delivered to this client" in captured.out
    assert "no leads were delivered" in captured.err
    assert len(ledger_items(ws["db"])) == 2                         # an empty delivery records nothing


def test_dry_run_previews_without_recording_anything(ws, capsys):
    code, preview = deliver(ws, "preview", "--dry-run")
    out = capsys.readouterr().out
    assert code == 0
    files = client_files(preview, preview=True)
    assert set(files) == {"csv", "xlsx", "html"}
    assert client_files(preview) == {}                              # no real (sendable) files
    assert "PREVIEW" in out and "leadgen deliver --client demo-client" in out
    _, preview_rows = read_rows(files["csv"])
    assert len(preview_rows) == 10
    assert ledger_items(ws["db"]) == {}                             # nothing recorded as delivered

    # so the real delivery afterwards still gets the full ten
    code, real = deliver(ws, "real")
    assert code == 0
    _, rows = read_rows(client_files(real)["csv"])
    assert [r["Company"] for r in rows] == [r["Company"] for r in preview_rows]
    assert ws["http"].calls == []


# --- 5. outbound-only commands refuse in delivery mode ----------------------------------------------------

@pytest.mark.parametrize("argv", [
    ["replies", "-p", DELIVERY_DEMO, "--file", REPLIES],
    ["replies", "--file", REPLIES],                                 # no -p: the default mode (delivery)
    ["serve", "-p", DELIVERY_DEMO, "--port", "0"],
    ["serve", "-p", LIVE_DELIVERY, "--token", "t0ken"],
    ["followups"],
    ["followups", "-p", DELIVERY_DEMO, "--days", "30"],
    ["followups", "--done", "1"],
], ids=["replies-delivery", "replies-no-playbook", "serve-demo", "serve-live", "followups",
        "followups-delivery", "followups-done"])
def test_outbound_only_commands_refuse_in_delivery_mode(ws, capsys, monkeypatch, argv):
    def boom(*a: Any, **k: Any) -> Any:
        raise AssertionError("an outbound-only command must not open anything in delivery mode")

    import leadgen.server
    monkeypatch.setattr(cli, "open_session", boom)
    monkeypatch.setattr(cli, "open_store", boom)
    monkeypatch.setattr(leadgen.server, "make_server", boom)
    code = main([*argv, *ws["g"]])
    err = capsys.readouterr().err
    assert code == 2, err
    assert "outbound-mode feature" in err and "mode: outbound" in err
    assert not ws["db"].exists()
    assert ws["http"].calls == []


# --- 6. outbound mode still works end to end --------------------------------------------------------------

def test_outbound_demo_still_runs_end_to_end(ws, capsys):
    out_root = ws["tmp"] / "outbound"
    assert main(["run", "-p", OUTBOUND_DEMO, "--out", str(out_root), *ws["g"]]) == 0
    (run_dir,) = [p for p in out_root.iterdir() if p.is_dir()]
    summary = json.loads((run_dir / "summary.json").read_text(encoding="utf-8"))
    counts = summary["counts"]
    assert summary["errors"] == []
    assert counts["written"] > 0 and counts["exported"] > 0

    # sequences were written and handed over in the sending tools' upload formats
    with open(run_dir / "instantly_upload.csv", newline="", encoding="utf-8") as f:
        uploads = list(csv.DictReader(f))
    assert len(uploads) == counts["exported"]
    for u in uploads:
        assert u["email"] and u["subject_1"]
        assert all(u[f"email_{i}"].strip() for i in range(1, 5)), u["email"]
    assert (run_dir / "smartlead_upload.csv").is_file() and (run_dir / "leads.json").is_file()
    with open(run_dir / "opportunities.csv", newline="", encoding="utf-8") as f:
        written = [r for r in csv.DictReader(f) if r["email_1"].strip()]
    assert len(written) == counts["written"]

    # the outbound-only commands work with an outbound playbook
    assert main(["replies", "-p", OUTBOUND_DEMO, "--file", REPLIES, "--out", str(ws["tmp"] / "replies"),
                 *ws["g"]]) == 0
    assert (ws["tmp"] / "replies" / "replies_classified.csv").is_file()
    assert main(["followups", "-p", OUTBOUND_DEMO, "--days", "30", *ws["g"]]) == 0
    out = capsys.readouterr().out
    assert "follow-up" in out
    assert ws["http"].calls == []


# --- 7. the docs match the code ---------------------------------------------------------------------------

DOCS = [REPO / "README.md", REPO / "PLAN.md", REPO / "docs" / "ARCHITECTURE.md"]
_GLOBAL_WITH_VALUE = {"-p", "--playbook", "--db", "--env-file"}
_STOP = {"|", "||", "&&", ">>", ">", "#", ";", "2>&1"}


def _code_blocks(text: str) -> List[str]:
    return re.findall(r"^```[^\n]*\n(.*?)^```", text, flags=re.M | re.S)


def _documented_commands() -> List[Tuple[str, str, List[str]]]:
    """(doc file, line, tokens after `leadgen`) for every leadgen command line in a fenced code block."""
    found = []
    for doc in DOCS:
        for block in _code_blocks(doc.read_text(encoding="utf-8")):
            for line in block.splitlines():
                if re.match(r"\s*(from|import)\s", line):          # Python code, not a command
                    continue
                for m in re.finditer(r"(?:^|[\s/(`])leadgen\s+(.*)$", line):
                    tokens: List[str] = []
                    for tok in m.group(1).split():
                        if tok in _STOP or tok.startswith("#"):
                            break
                        tokens.append(tok)
                    found.append((doc.name, line.strip(), tokens))
    return found


def _subparsers() -> Dict[str, Any]:
    parser = build_parser()
    for action in parser._actions:                                  # noqa: SLF001 - argparse has no public API
        if getattr(action, "choices", None) and isinstance(action.choices, dict):
            return dict(action.choices)
    raise AssertionError("no sub-commands found")


def _options(parser: Any) -> Set[str]:
    return {o for a in parser._actions for o in a.option_strings}   # noqa: SLF001


def test_every_documented_command_and_flag_exists():
    subs = _subparsers()
    top = _options(build_parser())
    commands = _documented_commands()
    assert len(commands) > 40, "the README should show plenty of real commands"
    problems = []
    for doc, line, tokens in commands:
        cmd, flags, skip_next = None, [], False
        for tok in tokens:
            if skip_next:
                skip_next = False
                continue
            clean = tok.strip("[](),:`'\"")
            if re.match(r"--?[A-Za-z]", clean):
                flag = clean.split("=", 1)[0]
                flags.append(flag)
                skip_next = cmd is None and flag in _GLOBAL_WITH_VALUE
            elif cmd is None:
                cmd = clean
        if cmd is not None and (cmd.startswith("<") or cmd.upper() == cmd):
            continue                                                # a placeholder like <command> / COMMAND
        allowed = top | (_options(subs[cmd]) if cmd in subs else set())
        if cmd is not None and cmd not in subs:
            problems.append(f"{doc}: unknown command {cmd!r} in: {line}")
        problems += [f"{doc}: unknown option {f!r} for {cmd or 'leadgen'} in: {line}"
                     for f in flags if f not in allowed]
    assert problems == []


def test_every_command_is_in_the_readme_command_table():
    readme = (REPO / "README.md").read_text(encoding="utf-8")
    table = readme[readme.index("## Commands"):]
    for name in _subparsers():
        assert f"`leadgen {name}" in table, f"README command table is missing `leadgen {name}`"


def _slug(heading: str) -> str:
    """GitHub's anchor for a heading: lower-case, punctuation dropped, spaces -> '-'."""
    text = re.sub(r"[^\w\- ]", "", heading.strip().lower())
    return text.replace(" ", "-")


def test_readme_internal_links_point_at_real_headings():
    readme = (REPO / "README.md").read_text(encoding="utf-8")
    prose = re.sub(r"^```.*?^```", "", readme, flags=re.M | re.S)   # headings inside code blocks don't count
    anchors = {_slug(h) for h in re.findall(r"^#{1,6} (.+)$", prose, flags=re.M)}
    links = set(re.findall(r"\]\(#([^)]+)\)", readme))
    assert links, "the README should link its sections"
    assert sorted(links - anchors) == []


def _env_example() -> Dict[str, str]:
    out = {}
    for line in (REPO / ".env.example").read_text(encoding="utf-8").splitlines():
        m = re.match(r"^([A-Z][A-Z0-9_]*)=(.*)$", line)
        if m:
            out[m.group(1)] = m.group(2).split("#", 1)[0].strip()
    return out


def test_env_example_lists_every_variable_the_code_and_playbooks_read():
    needed: Set[str] = {"ADZUNA_APP_ID", "GOOGLE_APPLICATION_CREDENTIALS", "GOOGLE_SERVICE_ACCOUNT_JSON",
                        "LEADGEN_EXPORT_WEBHOOK_URL", "LEADGEN_WEBHOOK_TOKEN", "LEADGEN_PLUGINS"}
    for kind in registry.KINDS:
        for name in registry.available(kind):
            cls = registry.resolve(kind, name)
            if not cls.__module__.startswith("leadgen."):           # a plugin some other test registered
                continue
            env_key = getattr(cls, "env_key", "")
            if env_key:
                needed.add(env_key)
    shipped = list((REPO / "playbooks").rglob("*.yaml")) + list((REPO / "clients").glob("*.yaml"))
    for path in shipped:
        for line in path.read_text(encoding="utf-8").splitlines():
            code = line.split(" #", 1)[0] if not line.lstrip().startswith("#") else ""
            needed.update(re.findall(r"\$\{([A-Z_][A-Z0-9_]*)", code))
    listed = _env_example()
    assert sorted(needed - set(listed)) == []
    # secrets ship empty, so nobody's real key is ever committed (only the SENDER_* examples have values)
    for name, value in listed.items():
        if not name.startswith("SENDER_"):
            assert value == "", f".env.example must not ship a value for {name}"
